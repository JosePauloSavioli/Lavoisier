#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import csv
import json
import re
from enum import Enum
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import ZipFile

from .ECS2_canonical_data import CANONICAL_KEY, ENERGY, INHERITANCE, PROCESS_TYPE, SPECIAL_ACTIVITY, TECHNOLOGY
from .utils import uuid_from_string, uuid_from_uuid


PROCESS_TYPE_REVERSE = {value.lower(): key for key, value in PROCESS_TYPE.items()}
SPECIAL_ACTIVITY_REVERSE = {value.lower(): key for key, value in SPECIAL_ACTIVITY.items()}
INHERITANCE_REVERSE = {value.lower(): key for key, value in INHERITANCE.items()}
TECHNOLOGY_REVERSE = {value.lower(): key for key, value in TECHNOLOGY.items()}
ENERGY_REVERSE = {value.lower(): key for key, value in ENERGY.items()}

PROCESS_IDENTITY = 'd1f70ac3-6914-4ba3-92b9-ffe7658c071d'
ACTIVITY_NAME_IDENTITY = 'c1df2e5f-15a7-433f-a7cb-9e41add21c57'
INTERMEDIATE_IDENTITY = '0f21f048-7d0d-43df-a54a-716397a7fba6'
ELEMENTARY_IDENTITY = '86ff6198-5381-46ec-b250-c3fb405b3303'
PARAMETER_IDENTITY = 'f20ad8f9-ed0e-4f66-8329-033ea94a1ce9'
PROPERTY_IDENTITY = '744c9fd5-db05-4f96-ae49-0371dfdeda1f'

PROCESS_LCA_TYPE = '6cd90a91-83d9-4c97-b4ec-aac59c8553ea'
SPECIAL_ACTIVITY = 'da8735d7-e473-4c5b-8532-25bc40aba761'
INHERITANCE_DEPTH = '85194c11-6de9-4272-a7ff-c64aa0bea6ea'
TECHNOLOGY_LEVEL = 'ffb442e6-9601-417e-b8f5-08061ed41441'
ENERGY_VALUE_BASIS = 'b97e1868-4c04-4d7e-bd0c-45c79ddc8299'
VALID_ENTIRE_PERIOD = '7e79e7af-71de-41a3-9dad-862e37cbe9ca'
MACRO_SCENARIO = 'a4146cc4-35d8-4a5e-8c75-9d3b717b4a24'
ALLOCATION_MODEL = '9767911c-ab5d-4eff-b150-215ca2a2ecb9'
GEOGRAPHICAL_SCOPE = '99a6248f-dbb0-49dc-a4ad-23abab6ba25f'
ECOINVENT_GEOGRAPHY = '107b99d0-ba7e-4825-b1c7-e4a25f2bfd08'
ECOINVENT_ROW = '51d34877-ad44-47b9-b157-12b8b05f8c3b'
POLITICAL_BOUNDARY = 'b629e5bf-48a8-4410-9103-f5cda1d9f729'
FLOW_LCA_TYPE = '5cc54527-0cb9-4180-82cc-d13240e94e50'
FLOW_IO_TYPE = '42b3c201-539f-4891-97cc-66156ac53083'
IS_REFERENCE = 'eb06d312-7738-42b2-a6c9-ac1fe9a018f4'
ELEMENTARY_DIRECTION = '735f3642-52b0-4f1c-90b3-78b0f7bafbeb'
ELEMENTARY_CONTEXT = '5119af08-62a1-4cb3-aa30-527cd72b5d90'
MAJOR_CONTEXT = '35478f20-e60e-4aad-82a9-1fcf0f4aea98'
AIR_HEIGHT = 'f84fa310-200d-4359-bc12-7b99c4659ec6'
AIR_SETTLEMENT = '9398d975-d406-4276-9755-8c01804a9d61'
SOIL_CONTEXT = '67bba357-b663-452c-874e-4f0c063a3647'
WATER_CONTEXT = '8868d6d3-6776-443b-90eb-c9ad9119348b'
RESOURCE_CONTEXT = '5027f4e3-6af1-47c9-b151-12088a5200b0'
ECOINVENT_CONTEXT = '51d9968a-d5d3-41e7-9fcd-f127f54ab761'
PRODUCTION_VOLUME = 'b3b96e47-db1f-496c-b73d-55b91716a5e7'
PEDIGREE = '63435e58-cd43-4c7b-a916-b49bdf2aef49'
ADDITIONAL_VARIANCE = '24964ee1-b167-4921-90d9-1b7aa18394ae'
VARIANCE_PROPERTY = '6cb57f90-10b6-406c-8d19-ac777ec074fa'
CAS_NUMBER = '55e46169-4f58-4204-a574-20703fdc02c4'
CHEMICAL_FORMULA = 'ed3b8a28-02a1-4a78-95d5-b6b1c3910d23'
CPC = 'c3988d68-211f-49f4-85c8-cbc47b30f83b'
ISIC = 'a158ab98-408a-430e-9cba-4bce9b85b9e1'
BY_PRODUCT = 'd695f836-5059-4e63-b6f0-bd093bc0e7f8'

UN_M49_REGISTRY = 'd11d8db1-8e7d-4041-baff-7ebd0d2fca68'
ECOINVENT_CONTEXT_ID = 'de659012-50c4-4e96-b54a-fc781bf987ab'
DEFAULT_SYSTEM_MODEL_ID = '8b738ea0-f89e-4627-8679-433616064e82'
DEFAULT_MACRO_ID = 'd9f57f0a-a01f-42eb-a57b-8f18d6635801'

STANDARD_PROPERTY_TO_ECS2 = {
    '3a0af1d6-04c3-41c6-a3da-92c4f61e0eaa': '3a0af1d6-04c3-41c6-a3da-92c4f61e0eaa',
    'd1a05915-1dd0-44b0-bc1b-f389e2ac3e52': '4d205f54-5f33-455c-a2bb-4650f1755e74',
    '6393c14b-db78-445d-a47b-c0cb866a1b25': '6393c14b-db78-445d-a47b-c0cb866a1b25',
    '67f102e2-9cb6-4d20-aa16-bf74d8a03326': '67f102e2-9cb6-4d20-aa16-bf74d8a03326',
    '5ac08789-7f70-439a-b60e-18c0d612d4ac': '6bc8ac5e-cc4c-45fb-93dc-53c0e45091df',
    '6d9e1462-80e3-4f10-b3f4-71febd6f1168': '6d9e1462-80e3-4f10-b3f4-71febd6f1168',
    'a9358458-9724-4f03-b622-106eda248916': 'a9358458-9724-4f03-b622-106eda248916',
    'c74c3729-e577-4081-b572-a283d2561a75': 'c74c3729-e577-4081-b572-a283d2561a75',
    '2aa79563-670b-448a-9edf-28f4954dfc74': 'e40445b1-8aba-4282-851d-9c2d28b1a5c1',
}

STANDARD_PROPERTY_NAMES = {
    '3a0af1d6-04c3-41c6-a3da-92c4f61e0eaa': 'dry mass',
    'd1a05915-1dd0-44b0-bc1b-f389e2ac3e52': 'capacity',
    '6393c14b-db78-445d-a47b-c0cb866a1b25': 'carbon content, non-fossil',
    '67f102e2-9cb6-4d20-aa16-bf74d8a03326': 'wet mass',
    '5ac08789-7f70-439a-b60e-18c0d612d4ac': 'energy content',
    '6d9e1462-80e3-4f10-b3f4-71febd6f1168': 'water in wet mass',
    'a9358458-9724-4f03-b622-106eda248916': 'water content',
    'c74c3729-e577-4081-b572-a283d2561a75': 'carbon content, fossil',
    '2aa79563-670b-448a-9edf-28f4954dfc74': 'density',
}

ACCESS_TO_ECS2 = {
    'public': 0,
    'licensees': 1,
    'restricted to licensed ecoinvent users': 1,
    'results only': 2,
    'restricted': 3,
}

DATA_PUBLISHED_IN_REVERSE = {
    'data as such not published elsewhere': 0,
    'data of some unit processes or subsystems are published': 1,
    'data published entirely in this source': 2,
}

NO_MATCH_TYPES = {
    'NO_MATCH_MAPPING',
    'NO MATCH',
    'NO_FLOW_MATCH_MANUAL',
    'NO_MAPPING',
    'NO MATCH MAPPING',
}

NON_ECS2_PROPERTIES = {
    PRODUCTION_VOLUME,
    PEDIGREE,
    ADDITIONAL_VARIANCE,
    VARIANCE_PROPERTY,
}

DIRECT_CATEGORICAL_PROPERTIES = {
    PROCESS_IDENTITY,
    ACTIVITY_NAME_IDENTITY,
    INTERMEDIATE_IDENTITY,
    ELEMENTARY_IDENTITY,
    PARAMETER_IDENTITY,
    PROPERTY_IDENTITY,
    PROCESS_LCA_TYPE,
    SPECIAL_ACTIVITY,
    INHERITANCE_DEPTH,
    TECHNOLOGY_LEVEL,
    ENERGY_VALUE_BASIS,
    VALID_ENTIRE_PERIOD,
    MACRO_SCENARIO,
    ALLOCATION_MODEL,
    GEOGRAPHICAL_SCOPE,
    ECOINVENT_GEOGRAPHY,
    ECOINVENT_ROW,
    POLITICAL_BOUNDARY,
    FLOW_LCA_TYPE,
    FLOW_IO_TYPE,
    IS_REFERENCE,
    ELEMENTARY_DIRECTION,
    ELEMENTARY_CONTEXT,
    MAJOR_CONTEXT,
    AIR_HEIGHT,
    AIR_SETTLEMENT,
    SOIL_CONTEXT,
    WATER_CONTEXT,
    RESOURCE_CONTEXT,
    ECOINVENT_CONTEXT,
    PEDIGREE,
    CAS_NUMBER,
    CHEMICAL_FORMULA,
    CPC,
    ISIC,
    BY_PRODUCT,
}


class CanonicalToECS2:

    def __init__(self, source, target, config=None):
        self.source = source
        self.target_type = target if isinstance(target, type) else type(target)
        self.config = config or {}
        self.language = self.config.get('default_language', 'en')
        self.project = None
        self.processes = {}
        self.production_systems = {}
        self.parameterizations = {}
        self.category_systems = {}
        self.units = {}
        self.flows = {}
        self.sources = {}
        self.actors = {}
        self.categories = {}
        self.property_registry = {}
        self._ucum_instance = None
        self._ucum_failed = False
        self._master_seen = {}
        self._classification_masters = {}
        self._entry_variables = {}
        self._property_metadata_cache = {}
        self._current_process = None
        self._current_system = None
        self._current_target = None
        self._current_log_key = None
        self._current_parameterizations = []
        self._logged_losses = set()
        self._load_bundle()
        self._index_registries()
        self._index_parameters()
        self.elementary_mapping = self._load_elementary_mapping(
            self.config.get('elementary_flow_mapping'),
            Path(__file__).parents[1] / 'Mappings' / 'ilcd1_to_ecs2_elementary_flows.json',
        )
        self.compartment_mapping = self._load_mapping(
            self.config.get('compartment_mapping'),
            Path(__file__).parents[1] / 'Mappings' / 'ilcd1_to_ecs2_compartments.json',
        )
        self.geography_mapping = self._load_mapping(
            self.config.get('geography_mapping'),
            Path(__file__).parents[1] / 'Mappings' / 'ilcd1_to_ecs2_geographies.json',
        )
        self.unit_mapping = self._load_mapping(self.config.get('unit_mapping'))
        self.log = {
            'summary': {
                'activities': 0,
                'warnings': 0,
                'unconverted': 0,
                'formulaMaterializations': 0,
                'mappedElementaryFlows': 0,
            },
            'global': [],
            'activities': {},
        }

    def convert(self):
        self._check_exportable()
        materializations = self._materializations()
        outputs = []
        for process, system in materializations:
            outputs.append(self._convert_activity(process, system))
        self.log['summary']['activities'] = len(outputs)
        self._current_log_key = None
        self._current_process = None
        self._current_system = None
        self._log_orphan_files(materializations)
        self._log_bundle_losses()
        return outputs, self.log

    def _load_bundle(self):
        values = []
        source = self.source
        if isinstance(source, (str, Path)):
            source = self._read_json(source)
        if isinstance(source, dict) and not source.get('datasetType'):
            for key in ('project', 'projects', 'processes', 'productionSystems', 'parameterizations', 'categorySystems', 'models'):
                value = source.get(key)
                if isinstance(value, dict) and value.get('datasetType'):
                    values.append(value)
                elif isinstance(value, dict):
                    values.extend(value.values())
                elif isinstance(value, (list, tuple)):
                    values.extend(value)
                elif value is not None:
                    values.append(value)
        elif isinstance(source, (list, tuple)):
            values.extend(source)
        else:
            values.append(source)

        expanded = []
        for value in values:
            if isinstance(value, (str, Path)):
                expanded.append(self._read_json(value))
            else:
                expanded.append(value)

        for value in expanded:
            dataset_type = self._get(value, 'datasetType')
            id_ = self._get(value, 'id')
            if dataset_type == 'project':
                if self.project is not None:
                    raise ValueError('ECS2 conversion accepts one canonical project dataset')
                self.project = value
            elif dataset_type == 'process':
                self.processes[id_] = value
            elif dataset_type == 'productionSystem':
                self.production_systems[id_] = value
            elif dataset_type == 'parameterization':
                self.parameterizations[id_] = value
            elif dataset_type == 'categorySystem':
                self.category_systems[id_] = value
            elif value is not None:
                raise ValueError(f'Unknown canonical datasetType: {dataset_type!r}')

        if not self.processes:
            raise ValueError('ECS2 conversion requires at least one canonical process dataset')

    def _read_json(self, path):
        path = Path(path)
        if path.suffix.lower() == '.zip':
            with ZipFile(path) as archive:
                return [
                    json.loads(archive.read(name))
                    for name in archive.namelist()
                    if name.endswith('.json')
                ]
        return json.loads(path.read_text())

    def _index_registries(self):
        datasets = list(self.processes.values()) + list(self.production_systems.values()) + list(self.parameterizations.values())
        if self.project is not None:
            datasets.append(self.project)
        for dataset in datasets:
            registry = self._get(dataset, 'registry')
            for name, target in (
                ('units', self.units),
                ('flows', self.flows),
                ('sources', self.sources),
                ('actors', self.actors),
                ('categories', self.categories),
                ('properties', self.property_registry),
            ):
                for item in self._get(registry, name, []) or []:
                    id_ = self._get(item, 'id') or (self._get(item, 'fileId') if name == 'categories' else None)
                    if id_:
                        target[id_] = item

    def _index_parameters(self):
        self._global_parameter_index = {}
        for parameterization in self.parameterizations.values():
            dataset_id = self._get(parameterization, 'id')
            for model in self._get(parameterization, 'models', []) or []:
                for parameter in self._get(model, 'parameters', []) or []:
                    parameter_id = self._get(parameter, 'id')
                    if not parameter_id:
                        continue
                    self._global_parameter_index[(dataset_id, parameter_id)] = parameter
                    self._global_parameter_index.setdefault(parameter_id, parameter)

    def _load_mapping(self, value, default=None):
        if value is None and default is not None and default.exists():
            value = default
        if value is None:
            return {}
        if isinstance(value, dict):
            return value
        if isinstance(value, list):
            return value
        path = Path(value)
        suffix = path.suffix.lower()
        if suffix == '.json':
            return json.loads(path.read_text())
        if suffix == '.csv':
            with path.open(newline='', encoding='utf-8-sig') as stream:
                return list(csv.DictReader(stream))
        if suffix in ('.xlsx', '.xlsm'):
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() if value is not None else '' for value in next(rows)]
            return [
                {header: value for header, value in zip(headers, row) if header}
                for row in rows
                if any(value is not None for value in row)
            ]
        raise ValueError(f'Unsupported mapping file: {path}')

    def _load_elementary_mapping(self, value, default=None):
        mapping = self._load_mapping(value, default)
        indexed = {}
        rows = mapping.items() if isinstance(mapping, dict) else ((None, row) for row in mapping)
        for key, row in rows:
            source_id = self._row(
                row,
                'SourceFlowUUID',
                'SourceFlowId',
                'sourceFlowUUID',
                'source_uuid',
                'sourceId',
            ) or key
            if source_id:
                indexed[str(source_id).strip().lower()] = row
        return indexed

    def _row(self, row, *names):
        if not isinstance(row, dict):
            return None
        normalized = {re.sub(r'[^a-z0-9]', '', str(key).lower()): value for key, value in row.items()}
        for name in names:
            key = re.sub(r'[^a-z0-9]', '', name.lower())
            if key in normalized and normalized[key] not in (None, ''):
                return normalized[key]
        return None

    def _check_exportable(self):
        def scalar(value, path):
            quantity_type = self._get(value, 'quantityType')
            category_type = self._get(value, 'categoryType')
            if quantity_type == 'resultingQuantity':
                if self._get(self._resolve_scalar_entry(value), 'quantityType') != 'singleQuantity':
                    raise ValueError(f'ECS2 export cannot resolve scalar resulting quantity at {path}')
            elif quantity_type and quantity_type != 'singleQuantity':
                raise ValueError(f'ECS2 export requires solved scalar quantities; {path} has {quantity_type}')
            if category_type == 'resultingCategory':
                if not self._is_scalar_category(self._resolve_scalar_entry(value)):
                    raise ValueError(f'ECS2 export cannot resolve scalar resulting category at {path}')
            elif category_type and category_type != 'singleCategory':
                raise ValueError(f'ECS2 export requires solved scalar categories; {path} has {category_type}')

        def walk(value, path='canonical'):
            scalar(value, path)
            if isinstance(value, dict):
                for key, item in value.items():
                    walk(item, f'{path}.{key}')
            elif isinstance(value, (list, tuple)):
                for index, item in enumerate(value):
                    walk(item, f'{path}[{index}]')
            else:
                for field in getattr(type(value), 'model_fields', {}):
                    walk(getattr(value, field), f'{path}.{field}')

        for parameterization in self.parameterizations.values():
            for model in self._get(parameterization, 'models', []) or []:
                for parameter in self._get(model, 'parameters', []) or []:
                    parameter_type = self._get(parameter, 'parameterType')
                    if parameter_type == 'quantitySetInput':
                        raise ValueError(
                            f'ECS2 export requires solved scalar parameterization; parameter {self._get(parameter, "id")} is a quantitySetInput'
                        )
                    if parameter_type == 'formula' and self._get(parameter, 'resultingEntry') is None:
                        raise ValueError(
                            f'ECS2 export requires solved formulas; parameter {self._get(parameter, "id")} has no resultingEntry'
                        )

        datasets = (
            [('process', value) for value in self.processes.values()]
            + [('productionSystem', value) for value in self.production_systems.values()]
            + [('parameterization', value) for value in self.parameterizations.values()]
            + ([('project', self.project)] if self.project is not None else [])
        )
        for type_, dataset in datasets:
            walk(dataset, f'{type_}[{self._get(dataset, "id")}]')

    def _materializations(self):
        occurrences = {id_: [] for id_ in self.processes}
        materializations = []
        for system in self.production_systems.values():
            for reference in self._get(system, 'processes', []) or []:
                process_id = self._reference_file(reference)
                process = self.processes.get(process_id)
                if process is None:
                    self._log('warning', 'missing_process', f'Production system references unavailable process {process_id}', systemId=self._get(system, 'id'))
                    continue
                occurrences[process_id].append(system)
                materializations.append((process, system))
        for process_id, process in self.processes.items():
            if not occurrences[process_id]:
                materializations.append((process, None))
        self._process_occurrences = occurrences
        return materializations

    def _convert_activity(self, process, system):
        self._current_process = process
        self._current_system = system
        self._current_target = self.target_type()
        self._master_seen = {}
        self._classification_masters = {}
        self._entry_variables = {}
        self._property_metadata_cache = {}
        self._current_parameterizations = self._parameterizations_for(process, system)
        materialized_id = self._materialized_process_id(process, system)
        self._current_log_key = materialized_id
        self.log['activities'][materialized_id] = {
            'processId': self._get(process, 'id'),
            'productionSystemId': self._get(system, 'id') if system is not None else None,
            'messages': [],
        }
        self._build_formula_environment(process, self._current_parameterizations)

        props = self._props(process)
        activity = self._activity(process, system, materialized_id, props)
        geography = self._geography(process, props)
        time_period = self._time(process, props)
        macro = self._macro_scenario(process, props)
        system_model = self._system_model(system)

        self._current_target.main_activity_type = 'childActivityDataset' if activity['inheritanceDepth'] else 'activityDataset'
        self._current_target.activityDescription = type(self._current_target.activityDescription)({
            'activity': activity,
            **({'classification': classifications} if (classifications := self._classifications(process, props, 1)) else {}),
            'geography': geography,
            'technology': self._technology(props),
            'timePeriod': time_period,
            'macroEconomicScenario': macro,
        })
        self._current_target.flowData = type(self._current_target.flowData)(self._flow_data(process, system))
        self._current_target.modellingAndValidation = type(self._current_target.modellingAndValidation)(
            self._modelling(system, system_model)
        )
        self._current_target.administrativeInformation = type(self._current_target.administrativeInformation)(
            self._administrative(process)
        )
        self._activity_index(activity, geography, time_period, system_model)

        self._log(
            'info',
            'activity_created',
            f'Created EcoSpold2 activity from canonical process {self._get(process, "id")}',
            exchanges=len(self._get(process, 'exchanges', []) or []),
            parameterizations=len(self._current_parameterizations),
        )
        return self._current_target

    def _materialized_process_id(self, process, system):
        process_id = self._get(process, 'id')
        occurrences = self._process_occurrences.get(process_id, [])
        if system is None or len(occurrences) <= 1:
            return process_id
        return uuid_from_string(f'canonical ECS2 materialization:{process_id}:{self._get(system, "id")}')

    def _parameterizations_for(self, process, system):
        ids = []
        if system is not None:
            for reference in self._get(system, 'parameterizations', []) or []:
                id_ = self._reference_file(reference)
                if id_ and id_ not in ids:
                    ids.append(id_)

        for id_ in self._result_parameterization_ids(process):
            if id_ not in ids:
                ids.append(id_)

        result = []
        for id_ in ids:
            if id_ in self.parameterizations:
                result.append(self.parameterizations[id_])
            else:
                self._log(
                    'warning',
                    'missing_parameterization',
                    f'Canonical scalar result references unavailable parameterization {id_}',
                    systemId=self._get(system, 'id') if system is not None else None,
                    processId=self._get(process, 'id'),
                )
        return result

    def _result_parameterization_ids(self, value):
        result = set()

        def walk(entry):
            quantity_type = self._get(entry, 'quantityType')
            category_type = self._get(entry, 'categoryType')
            if quantity_type == 'resultingQuantity' or category_type == 'resultingCategory':
                reference = self._get(entry, 'parameter')
                file_id = self._get(reference, 'fileId')
                if file_id:
                    result.add(file_id)

            if isinstance(entry, dict):
                for child in entry.values():
                    walk(child)
            elif isinstance(entry, (list, tuple)):
                for child in entry:
                    walk(child)
            else:
                fields = getattr(type(entry), 'model_fields', None)
                if fields:
                    for field in fields:
                        walk(getattr(entry, field))

        walk(value)
        return result

    def _activity(self, process, system, materialized_id, props):
        explicit_process = self._identity_entry(process, PROCESS_IDENTITY)
        process_id = self._ecs2_id(materialized_id)
        if explicit_process and materialized_id == self._get(process, 'id'):
            process_id = self._identity_id(self._get(process, 'id'), explicit_process, 'process')

        activity_identity = self._identity_entry(process, ACTIVITY_NAME_IDENTITY)
        activity_identity_value = self._resolve_scalar_entry(self._get(activity_identity, 'value'))
        activity_name = (
            self._get(activity_identity_value, 'label')
            or self._category_scalar(activity_identity_value)
            or self._get(process, 'name')
            or 'Unnamed canonical process'
        )
        if activity_identity:
            activity_name_id = self._get(activity_identity_value, 'entryId')
            if not activity_name_id:
                activity_name_id = self._ecs2_id(uuid_from_string(f'canonical activity name:{activity_name}'))
        else:
            activity_name_id = self._ecs2_id(uuid_from_string(f'canonical activity name:{activity_name}'))

        process_type = self._category_scalar(self._property_value(props, PROCESS_LCA_TYPE))
        special = self._category_scalar(self._property_value(props, SPECIAL_ACTIVITY))
        inheritance = self._category_scalar(self._property_value(props, INHERITANCE_DEPTH))
        energy = self._category_scalar(self._property_value(props, ENERGY_VALUE_BASIS))
        inheritance_code = self._reverse_enum(inheritance, INHERITANCE_REVERSE, 0)

        activity = {
            'id': process_id,
            'activityNameId': activity_name_id,
            'inheritanceDepth': inheritance_code,
            'type': self._reverse_enum(process_type, PROCESS_TYPE_REVERSE, 1),
            'specialActivityType': self._reverse_enum(special, SPECIAL_ACTIVITY_REVERSE, 0),
            'activityName': self._text(activity_name),
        }
        if energy is not None:
            activity['energyValues'] = self._reverse_enum(energy, ENERGY_REVERSE, 0)

        additional = self._get(process, 'processAdditionalInformation')
        included = self._get(additional, 'includedProcesses', []) or []
        if included:
            entries = []
            for item in included:
                name = self._get(item, 'processName')
                description = self._get(item, 'description')
                if len(included) == 1 and name == activity_name and description:
                    entries.append(description)
                elif description:
                    entries.append(f'{name}: {description}' if name else description)
                elif name:
                    entries.append(name)
            if entries:
                activity['includedActivitiesEnd'] = self._texts(entries)

        comments = self._context_comments(process, system)
        if comments:
            activity['generalComment'] = {'text': self._texts(comments)}

        parent = self._dynamic_identity(props, 'ecoinvent parent activity id')
        if parent:
            activity['parentActivityId'] = parent
        allocation = self._dynamic_identity(props, 'ecoinvent master allocation property id')
        if allocation:
            activity['masterAllocationPropertyId'] = allocation
        return activity

    def _context_comments(self, process, system):
        comments = self._comment_texts(self._get(process, 'comments'))
        comments.extend(self._flattened_category_comments(process))
        additional = self._get(process, 'processAdditionalInformation')
        for value in self._get(additional, 'additionalProcessDescription', []) or []:
            comments.append(f'Additional process description: {value}')
        if self._get(additional, 'timeSeriesStep') is not None:
            step = self._get(additional, 'timeSeriesStep')
            comments.append(f'Canonical time-series step: {self._quantity_amount(step)} {self._unit_name(self._get(step, "unit"))}')

        if system is not None:
            for text in self._comment_texts(self._get(system, 'comments')):
                comments.append(f'Production system: {text}')
            comments.extend(f'Production system: {text}' for text in self._flattened_category_comments(system))

        for parameterization in self._current_parameterizations:
            for text in self._comment_texts(self._get(parameterization, 'comments')):
                comments.append(f'Parameterization: {text}')

        if self.project is not None:
            description = self._get(self.project, 'description')
            intended = self._get(self.project, 'intendedApplication')
            handling = self._get(self._get(self.project, 'modelling'), 'dataHandlingComment')
            if description:
                comments.append(f'Project: {description}')
            if intended:
                comments.append(f'Intended application: {intended}')
            if handling:
                comments.append(f'Data handling: {handling}')
            comments.extend(f'Project: {text}' for text in self._flattened_category_comments(self.project))

        for dataset, label in self._file_contexts(process):
            info = self._get(dataset, 'fileInformation')
            uri = self._get(info, 'downloadableURI')
            if uri and dataset is not process:
                comments.append(f'Canonical {label} source: {uri}')
        return comments

    def _unit_name(self, unit_id):
        unit = self.units.get(unit_id)
        return self._get(unit, 'shortName') or self._get(unit, 'semanticName') or unit_id or ''

    def _geography(self, process, props):
        scope = self._property(props, GEOGRAPHICAL_SCOPE)
        nested = self._get(scope, 'properties', []) or []
        direct = self._property(nested, ECOINVENT_GEOGRAPHY)
        direct_value = self._resolve_scalar_entry(self._get(direct, 'value'))
        entry_id = self._get(direct_value, 'entryId')
        shortname = self._get(direct_value, 'code')
        name = self._category_scalar(direct_value)
        m49 = None
        political = self._property(nested, POLITICAL_BOUNDARY)
        if political:
            political_value = self._resolve_scalar_entry(self._get(political, 'value'))
            if self._get(political, 'categorySystem') == UN_M49_REGISTRY:
                m49 = self._get(political_value, 'code')
                name = name or self._category_scalar(political_value)

        generated = False
        if entry_id:
            geography_id = entry_id
            shortname = shortname or self._geography_shortname(m49, name) or str(name or 'GLO')
        else:
            shortname = self._geography_shortname(m49, name)
            mapped = self._mapping_lookup(self.geography_mapping, m49, shortname, name)
            geography_id, mapped_name, mapped_short = self._geography_mapping_value(mapped)
            if geography_id:
                name = mapped_name or name
                shortname = mapped_short or shortname
            else:
                default = self.config.get('default_geography', 'GLO')
                mapped = self._mapping_lookup(self.geography_mapping, default)
                geography_id, mapped_name, mapped_short = self._geography_mapping_value(mapped)
                if not geography_id and default == 'GLO':
                    geography_id = self._ecs2_id(uuid_from_string('ecoinvent geography:GLO'))
                    generated = True
                shortname = mapped_short or default
                name = mapped_name or ('Global' if default == 'GLO' else default)
                self._log('warning', 'geography_defaulted', f'No EcoSpold2 geography mapping for process {self._get(process, "id")}; used {shortname}', m49=m49)

        if generated:
            self._master('geography', {
                'id': geography_id,
                'name': self._text(name or shortname),
                'shortname': self._text(shortname),
                **({'uNCode': int(m49)} if m49 and str(m49).isdigit() else {}),
            })
        return {'geographyId': geography_id, 'shortname': self._text(shortname)}

    def _geography_shortname(self, m49, name):
        if m49:
            code = str(m49).zfill(3)
            if code == '001':
                return 'GLO'
            try:
                import pycountry
                country = pycountry.countries.get(numeric=code)
                if country:
                    return country.alpha_2
            except Exception:
                pass
        if name and str(name).upper() in self.geography_mapping:
            return str(name).upper()
        return None

    def _geography_mapping_value(self, value):
        if value is None:
            return None, None, None
        if isinstance(value, str):
            return value, None, None
        if isinstance(value, dict):
            return (
                self._row(value, 'TargetGeographyUUID', 'geographyId', 'targetId', 'id'),
                self._row(value, 'TargetGeographyName', 'name'),
                self._row(value, 'TargetGeographyCode', 'shortname', 'shortName', 'code'),
            )
        return None, None, None

    def _technology(self, props):
        value = self._category_scalar(self._property_value(props, TECHNOLOGY_LEVEL))
        return {'technologyLevel': self._reverse_enum(value, TECHNOLOGY_REVERSE, 3)}

    def _time(self, process, props):
        additional = self._get(process, 'processAdditionalInformation')
        validity = self._get(additional, 'validity')
        start = self._serial(self._get(validity, 'from'))
        end = self._serial(self._get(validity, 'until'))
        if not start or not end:
            now = datetime.now(timezone.utc).date().isoformat()
            start = start or now
            end = end or now
            self._log('warning', 'time_defaulted', f'Process {self._get(process, "id")} lacks complete validity dates')
        valid = self._property_value(props, VALID_ENTIRE_PERIOD)
        return {
            'startDate': str(start)[:10],
            'endDate': str(end)[:10],
            'isDataValidForEntirePeriod': True if valid is None else bool(valid),
        }

    def _macro_scenario(self, process, props):
        prop = self._property(props, MACRO_SCENARIO)
        value = self._resolve_scalar_entry(self._get(prop, 'value'))
        entry_id = self._get(value, 'entryId')
        name = self._category_scalar(value)
        if not entry_id:
            entry_id = self.config.get('default_macro_scenario_id', DEFAULT_MACRO_ID)
            if not name:
                name = self.config.get('default_macro_scenario_name', 'Business-as-Usual')
            self._log('warning', 'macro_scenario_defaulted', f'Process {self._get(process, "id")} has no EcoSpold2 macro-economic scenario identity')
        return {'macroEconomicScenarioId': entry_id, 'name': self._text(name or 'Business-as-Usual')}

    def _classifications(self, owner, props, type_):
        result = []
        handled = set()
        for property_id, system_name in ((CPC, 'CPC'), (ISIC, 'ISIC Rev. 4'), (BY_PRODUCT, 'by-product classification')):
            for prop in self._properties(props, property_id):
                value = self._resolve_scalar_entry(self._get(prop, 'value'))
                entry_id = self._get(value, 'entryId')
                if not entry_id:
                    self._log('warning', 'classification_not_mapped', f'Classification {system_name} has no EcoSpold2 master-data identity', ownerId=self._get(owner, 'id'))
                    continue
                classification = self._classification(entry_id, system_name, value)
                result.append(classification)
                handled.add(self._get(prop, 'id'))

        for prop in props:
            if self._get(prop, 'id') in handled:
                continue
            registry = self.property_registry.get(self._get(prop, 'property'))
            name = self._get(registry, 'name') or ''
            match = re.fullmatch(r'ECS2 classification (.+)', name, re.IGNORECASE)
            if not match:
                continue
            value = self._resolve_scalar_entry(self._get(prop, 'value'))
            entry_id = self._get(value, 'entryId')
            if not entry_id:
                continue
            system_name = match.group(1)
            classification = self._classification(entry_id, system_name, value)
            result.append(classification)
        return result

    def _classification(self, entry_id, system_name, value):
        code = self._get(value, 'code')
        label = self._get(value, 'label')
        legacy_value = self._get(value, 'value')
        if legacy_value and label and not code and str(legacy_value) != str(label):
            code, label = label, legacy_value
        label = label or legacy_value or entry_id
        text = f'{code}: {label}' if code and str(code) != str(label) else str(label)
        return {
            'classificationId': entry_id,
            'classificationSystem': self._text(system_name),
            'classificationValue': self._text(text),
        }

    def _flattened_category_comments(self, owner):
        result = []
        for prop in self._props(owner):
            property_id = self._get(prop, 'property')
            category_system = self._get(prop, 'categorySystem')
            if (
                not category_system
                or property_id in DIRECT_CATEGORICAL_PROPERTIES
            ):
                continue
            registry = self.property_registry.get(property_id)
            name = self._get(registry, 'name') or 'Canonical category'
            if re.fullmatch(r'ECS2 classification .+', name, re.IGNORECASE):
                continue
            value = self._resolve_scalar_entry(self._get(prop, 'value'))
            label = self._category_scalar(value)
            if label is None:
                continue
            code = self._get(value, 'code')
            text = f'{code}: {label}' if code and str(code) != str(label) else str(label)
            result.append(f'{name}: {text}')
        return result

    def _classification_master(self, system_name, type_, classification):
        system_id = self._ecs2_id(uuid_from_string(f'ecoinvent classification system:{system_name}:{type_}'))
        value = {
            'id': classification['classificationId'],
            'name': classification['classificationValue'],
        }
        master = self._classification_masters.get(system_id)
        if master is None:
            master = self._current_target.userMaster.get_class('classificationSystem')()
            master.id = system_id
            master.type = type_
            master.name = self._text(system_name)
            self._classification_masters[system_id] = master
            setattr(self._current_target.userMaster, 'classificationSystem', master)
        seen = self._master_seen.setdefault('classificationValue', set())
        key = (system_id, classification['classificationId'])
        if key not in seen:
            master.classificationValue = value
            seen.add(key)

    def _flow_data(self, process, system):
        result = {'intermediateExchange': []}
        functional = self._functional_references(system)
        provider_links = self._provider_links(system, process)
        for exchange in self._get(process, 'exchanges', []) or []:
            flow = self.flows.get(self._get(exchange, 'flowable'))
            if flow is None:
                self._log('warning', 'missing_flow', f'Exchange {self._get(exchange, "id")} references unavailable flow {self._get(exchange, "flowable")}')
                continue
            flow_type = self._flow_type(exchange, flow)
            converted = self._exchange(exchange, flow, flow_type, functional, provider_links)
            if converted is None:
                continue
            key = 'elementaryExchange' if flow_type == 'elementary' else 'intermediateExchange'
            result.setdefault(key, []).append(converted)

        parameters = self._parameters(process)
        if parameters:
            result['parameter'] = parameters
        if not result['intermediateExchange']:
            self._log('warning', 'no_intermediate_exchange', 'EcoSpold2 activity has no intermediate exchange after conversion')
        return result

    def _flow_type(self, exchange, flow):
        value = self._category_scalar(self._property_value(self._props(exchange), FLOW_LCA_TYPE))
        if value in ('intermediate', 'elementary'):
            return value
        if self._identity_entry(flow, ELEMENTARY_IDENTITY) or self._property(self._props(exchange), ELEMENTARY_CONTEXT):
            return 'elementary'
        return 'intermediate'

    def _exchange(self, exchange, flow, flow_type, functional, provider_links):
        quantity = self._resolve_scalar_entry(self._get(exchange, 'quantity'))
        if self._get(quantity, 'quantityType') != 'singleQuantity':
            self._log('warning', 'exchange_not_scalar', f'Exchange {self._get(exchange, "id")} does not resolve to a single quantity')
            return None
        unit_id = self._get(quantity, 'unit')
        unit = self.units.get(unit_id)
        if not unit:
            self._log('warning', 'missing_unit', f'Exchange {self._get(exchange, "id")} references unavailable unit {unit_id}')
            return None

        target_flow = None
        factor = Decimal('1')
        target_unit_name = self._ecs2_unit_name(unit)
        target_unit_id = self._target_unit_id(unit, target_unit_name)
        compartment = None
        existing_flow_master = flow_type == 'elementary'

        if flow_type == 'elementary':
            target_flow, target_unit_name, factor, compartment = self._elementary_target(flow, exchange, quantity, unit)
            if target_flow is None:
                return None
            target_flow_id = target_flow['id']
            target_unit_id = self._target_unit_id(unit, target_unit_name)
        else:
            identity = self._identity_entry(flow, INTERMEDIATE_IDENTITY)
            existing_flow_master = bool(identity)
            target_flow_id = self._identity_id(self._get(flow, 'id'), identity, 'intermediate flow') if identity else self._ecs2_id(self._get(flow, 'id'))
            identity_value = self._resolve_scalar_entry(self._get(identity, 'value')) if identity else None
            target_flow = {
                'id': target_flow_id,
                'name': self._get(identity_value, 'label') or self._category_scalar(identity_value) or self._get(flow, 'name') or target_flow_id,
            }

        amount = Decimal(str(self._get(quantity, 'amount'))) * factor
        exchange_id = self._ecs2_id(self._get(exchange, 'id'))
        direction = self._category_scalar(self._property_value(self._props(exchange), FLOW_IO_TYPE))
        if direction not in ('input', 'output'):
            direction = 'input' if amount < 0 else 'output'
            self._log('warning', 'direction_inferred', f'Exchange {self._get(exchange, "id")} has no canonical input/output type; inferred {direction}')

        data = {
            'id': exchange_id,
            'name': self._text(target_flow['name']),
            'amount': amount,
            'unitId': target_unit_id,
            'unitName': self._text(target_unit_name),
        }
        entry_key = ('exchange', self._get(exchange, 'id'))
        variable = self._entry_variables.get(entry_key)
        if variable:
            data['variableName'] = variable
        self._entry_formula(entry_key, factor)

        comments = self._comment_texts(self._get(exchange, 'comments'))
        comments.extend(self._flattened_category_comments(exchange))
        comments.extend(self._flattened_category_comments(flow))
        if comments:
            data['comment'] = self._texts(comments)
        source = self._first_source(self._get(exchange, 'sourceIds'))
        if source:
            data.update(source)

        uncertainty = self._uncertainty(
            quantity,
            self._entry_metadata_props(entry_key, self._exchange_quantity_metadata(exchange)),
            factor,
            'exchange',
        )
        if uncertainty:
            data['uncertainty'] = uncertainty

        properties = self._exchange_properties(flow, exchange, factor)
        if properties:
            data['property'] = properties

        cas = self._category_scalar(self._property_value(self._props(flow), CAS_NUMBER))
        if cas:
            data['casNumber'] = cas

        if flow_type == 'elementary':
            data['elementaryExchangeId'] = target_flow_id
            data['compartment'] = compartment
            if direction == 'input':
                data['inputGroup'] = 4
            else:
                data['outputGroup'] = 4
            formula = self._category_scalar(self._property_value(self._props(flow), CHEMICAL_FORMULA))
            if formula:
                data['formula'] = formula
            if not existing_flow_master:
                self._elementary_master(target_flow, target_unit_id, target_unit_name, compartment, flow)
        else:
            data['intermediateExchangeId'] = target_flow_id
            if direction == 'input':
                data['inputGroup'] = 5
            else:
                data['outputGroup'] = 0 if self._get(exchange, 'id') in functional or self._bool_property(exchange, IS_REFERENCE) else 2
            link = provider_links.get(self._get(exchange, 'id'))
            if link:
                data['activityLinkId'] = link
            production = self._property(self._props(exchange), PRODUCTION_VOLUME)
            if production:
                production_quantity = self._resolve_scalar_entry(self._get(production, 'value'))
                production_unit = self.units.get(self._get(production_quantity, 'unit'))
                if production_unit:
                    production_factor = self._unit_factor(
                        self._get(production_unit, 'shortName'),
                        target_unit_name,
                        production_unit,
                        unit,
                    )
                    if production_factor is None:
                        self._log('warning', 'production_volume_unit', f'Could not convert production volume for flow {self._get(flow, "id")} to exchange unit {target_unit_name}')
                    else:
                        data['productionVolumeAmount'] = Decimal(str(self._get(production_quantity, 'amount'))) * production_factor
                        production_key = ('property', self._get(production, 'id'))
                        production_variable = self._entry_variables.get(production_key)
                        if production_variable:
                            data['productionVolumeVariableName'] = production_variable
                        self._entry_formula(production_key, production_factor)
                        production_comments = self._comment_texts(self._get(production, 'comments'))
                        if production_comments:
                            data['productionVolumeComment'] = self._texts(production_comments)
                        production_uncertainty = self._uncertainty(
                            production_quantity,
                            self._entry_metadata_props(
                                ('property', self._get(production, 'id')),
                                self._production_volume_metadata(exchange),
                            ),
                            production_factor,
                            'production volume',
                        )
                        if production_uncertainty:
                            data['productionVolumeUncertainty'] = production_uncertainty
                        source = self._first_source(self._get(production, 'sourceIds'))
                        if source:
                            for key, value in source.items():
                                data['productionVolume' + key[0].upper() + key[1:]] = value
            classifications = self._classifications(flow, self._props(flow), 2)
            if classifications:
                data['classification'] = classifications
            if not existing_flow_master:
                self._intermediate_master(target_flow, target_unit_id, target_unit_name, flow, classifications)
        return data

    def _elementary_target(self, flow, exchange, quantity, unit):
        identity = self._identity_entry(flow, ELEMENTARY_IDENTITY)
        if identity:
            target_id = self._identity_id(self._get(flow, 'id'), identity, 'elementary flow')
            target_name = self._category_scalar(self._get(identity, 'value')) or self._get(flow, 'name') or target_id
            context = self._canonical_compartment(exchange)
            if context is None:
                self._log('warning', 'missing_compartment', f'Elementary flow {self._get(flow, "id")} has EcoSpold2 identity but no convertible compartment')
                return None, None, None, None
            return {'id': target_id, 'name': target_name}, self._ecs2_unit_name(unit), Decimal('1'), context

        row = None
        for candidate in self._source_entry_ids(flow):
            row = self.elementary_mapping.get(str(candidate).lower())
            if row:
                break
        if row is None:
            self._log('warning', 'elementary_flow_unmapped', f'No GLAD elementary-flow mapping for canonical flow {self._get(flow, "id")}', flowName=self._get(flow, 'name'))
            self.log['summary']['unconverted'] += 1
            return None, None, None, None

        map_type = str(self._row(row, 'MapType', 'map_type') or '').strip().upper()
        target_id = self._row(row, 'TargetFlowUUID', 'targetFlowUUID', 'targetId')
        if map_type in NO_MATCH_TYPES or not target_id:
            self._log('warning', 'elementary_flow_no_match', f'GLAD mapping has no EcoSpold2 match for canonical flow {self._get(flow, "id")}', mapType=map_type)
            self.log['summary']['unconverted'] += 1
            return None, None, None, None

        target_name = self._row(row, 'TargetFlowName', 'targetName') or self._get(flow, 'name') or target_id
        source_unit = self._row(row, 'SourceUnit', 'sourceUnit')
        target_unit = self._row(row, 'TargetUnit', 'targetUnit') or source_unit or self._get(unit, 'shortName') or '1'
        conversion = self._decimal(self._row(row, 'ConversionFactor', 'conversionFactor'), Decimal('1'))
        canonical_unit = self._get(unit, 'shortName') or self._get(unit, 'semanticName') or '1'
        unit_factor = self._unit_factor(canonical_unit, source_unit or canonical_unit, unit)
        if unit_factor is None:
            self._log('warning', 'elementary_unit_unmapped', f'Cannot convert {canonical_unit} to GLAD source unit {source_unit} for flow {self._get(flow, "id")}')
            self.log['summary']['unconverted'] += 1
            return None, None, None, None
        factor = unit_factor * conversion

        context_name = self._row(row, 'TargetFlowContext', 'targetFlowContext')
        compartment = self._compartment_from_mapping(context_name)
        if compartment is None:
            self._log('warning', 'compartment_unmapped', f'No EcoSpold2 compartment mapping for GLAD context {context_name!r}', flowId=self._get(flow, 'id'))
            self.log['summary']['unconverted'] += 1
            return None, None, None, None

        self.log['summary']['mappedElementaryFlows'] += 1
        self._log('info', 'elementary_flow_mapped', f'Mapped canonical elementary flow {self._get(flow, "id")} to EcoSpold2 {target_id}', mapType=map_type or None)
        return {'id': target_id, 'name': target_name}, str(target_unit), factor, compartment

    def _source_entry_ids(self, flow):
        values = []
        for prop in self._props(flow):
            value = self._resolve_scalar_entry(self._get(prop, 'value'))
            entry_id = self._get(value, 'entryId')
            if entry_id:
                values.append(entry_id)
        values.append(self._get(flow, 'id'))
        return values

    def _canonical_compartment(self, exchange):
        composite = self._property(self._props(exchange), ELEMENTARY_CONTEXT)
        nested = self._get(composite, 'properties', []) or []
        ecoinvent = self._property(nested, ECOINVENT_CONTEXT)
        value = self._resolve_scalar_entry(self._get(ecoinvent, 'value'))
        entry_id = self._get(value, 'entryId')
        major = self._category_scalar(self._property_value(nested, MAJOR_CONTEXT))
        sub = None
        if major == 'air':
            height = self._category_scalar(self._property_value(nested, AIR_HEIGHT))
            settlement = self._category_scalar(self._property_value(nested, AIR_SETTLEMENT))
            sub = self._air_subcompartment(height, settlement)
        elif major == 'soil':
            sub = self._category_scalar(self._property_value(nested, SOIL_CONTEXT))
        elif major == 'water':
            sub = self._category_scalar(self._property_value(nested, WATER_CONTEXT))
        elif major == 'natural resource':
            sub = self._category_scalar(self._property_value(nested, RESOURCE_CONTEXT))

        label = self._category_scalar(value)
        if entry_id and label and ',' in str(label):
            major, sub = [part.strip() for part in str(label).split(',', 1)]
        elif (not major or not sub) and label and ',' in str(label):
            parts = [part.strip() for part in str(label).split(',', 1)]
            major = major or parts[0]
            sub = sub or parts[1]
        if entry_id and major and sub:
            result = {
                'subcompartmentId': entry_id,
                'compartment': self._text(major),
                'subcompartment': self._text(sub),
            }
            return result
        if major and sub:
            return self._compartment_from_mapping(f'{major}/{sub}')
        return None

    def _air_subcompartment(self, height, settlement):
        reverse = {
            ('high stacks', 'non-urban'): 'non-urban air or from high stacks',
            ('close to ground', 'urban'): 'urban air close to ground',
            (None, 'low population density'): 'low population density',
            (None, 'indoor'): 'indoor',
            ('lower stratosphere + upper troposphere', None): 'lower stratosphere + upper troposphere',
        }
        return reverse.get((height, settlement)) or settlement or height

    def _compartment_from_mapping(self, context_name):
        if not context_name:
            return None
        mapped = self._mapping_lookup(self.compartment_mapping, context_name)
        if isinstance(mapped, str):
            sub_id = mapped
        elif isinstance(mapped, dict):
            sub_id = self._row(mapped, 'subcompartmentId', 'id', 'TargetCompartmentUUID')
        else:
            return None
        parts = str(context_name).split('/', 1)
        if len(parts) != 2:
            return None
        result = {
            'subcompartmentId': sub_id,
            'compartment': self._text(parts[0]),
            'subcompartment': self._text(parts[1]),
        }
        return result

    def _compartment_master(self, compartment):
        major = self._text_value(compartment['compartment'])
        sub = self._text_value(compartment['subcompartment'])
        major_id = self._ecs2_id(uuid_from_string(f'ecoinvent compartment:{major}'))
        self._master('compartment', {
            'id': major_id,
            'name': self._text(major),
            'subcompartment': {
                'id': compartment['subcompartmentId'],
                'name': self._text(sub),
            },
        }, key=(major_id, compartment['subcompartmentId']))

    def _production_volume_metadata(self, exchange):
        exchange_id = self._get(exchange, 'id')
        ids = {
            uuid_from_string(f'ECS2:{exchange_id}:{PEDIGREE}:production volume:pedigree'),
            uuid_from_string(f'ECS2:{exchange_id}:{ADDITIONAL_VARIANCE}:production volume:pedigree-variance'),
            uuid_from_string(f'ECS2:{exchange_id}:{VARIANCE_PROPERTY}:production volume:variance'),
        }
        return [
            prop for prop in self._props(exchange)
            if self._get(prop, 'id') in ids
            or self._property_registry_name(self._get(prop, 'property')).lower().startswith(
                'ecs2 production volume uncertainty '
            )
        ]

    def _exchange_quantity_metadata(self, exchange):
        production = {self._get(prop, 'id') for prop in self._production_volume_metadata(exchange)}
        return [prop for prop in self._props(exchange) if self._get(prop, 'id') not in production]

    def _functional_references(self, system):
        result = set()
        if system is None:
            return result
        process_id = self._get(self._current_process, 'id')
        for item in self._get(system, 'functionalReference', []) or []:
            reference = self._get(item, 'reference')
            ref_process = self._reference_file(self._get(reference, 'processId'))
            if ref_process != process_id:
                continue
            exchange = self._get(reference, 'outputExchangeId')
            result.add(self._reference_entry(exchange, 'exchangeId'))
        return result

    def _provider_links(self, system, process):
        if system is None:
            return {}
        process_id = self._get(process, 'id')
        result = {}
        linking = self._get(system, 'linking') or {}
        foreground = self._get(self._get(linking, 'foregroundPartialLinkingMatrix'), 'linking', []) or []
        backgrounds = self._get(linking, 'backgroundPartialLinkingMatrices', []) or []
        rows = list(foreground)
        for matrix in backgrounds:
            rows.extend(self._get(matrix, 'linking', []) or [])

        for row in rows:
            consumer = self._get(row, 'id')
            consumer_process = self._reference_file(self._get(consumer, 'processId'))
            exchange_id = self._reference_entry(self._get(consumer, 'inputExchangeId'), 'exchangeId')
            if consumer_process != process_id or not exchange_id:
                continue
            values = self._get(row, 'values', []) or []
            if len(values) != 1 or Decimal(str(self._get(values[0], 'value'))) != Decimal('1'):
                self._log('warning', 'partial_link_not_representable', f'Exchange {exchange_id} has provider shares that cannot be represented by one EcoSpold2 activityLinkId', shares=len(values))
                continue
            provider = self._get(values[0], 'id')
            provider_process = self._resolve_scalar_entry(self._get(provider, 'processId'))
            if self._get(provider_process, 'categoryType') == 'singleCategory':
                entry_id = self._get(provider_process, 'entryId')
                if entry_id:
                    result[exchange_id] = entry_id
                    continue
                label = self._get(provider_process, 'label')
                if self._looks_uuid(label):
                    result[exchange_id] = label
                    continue
            provider_process_id = self._reference_file(provider_process)
            if provider_process_id and provider_process_id in self.processes:
                provider_identity = self._identity_entry(self.processes[provider_process_id], ACTIVITY_NAME_IDENTITY)
                if provider_identity:
                    result[exchange_id] = self._get(self._get(provider_identity, 'value'), 'entryId')
                else:
                    name = self._get(self.processes[provider_process_id], 'name') or provider_process_id
                    result[exchange_id] = self._ecs2_id(uuid_from_string(f'canonical activity name:{name}'))
        return result

    def _exchange_properties(self, flow, exchange, factor):
        result = []
        for owner in (flow, exchange):
            props = self._props(owner)
            for index, prop in enumerate(props):
                property_id = self._get(prop, 'property')
                if property_id in NON_ECS2_PROPERTIES:
                    continue

                value = self._resolve_scalar_entry(self._get(prop, 'value'))
                if self._get(value, 'quantityType') != 'singleQuantity':
                    continue

                source_unit = self.units.get(self._get(value, 'unit'))
                if not source_unit:
                    self._log('warning', 'property_unit_missing', f'Property {self._get(prop, "id")} references unavailable unit {self._get(value, "unit")}')
                    continue

                target_unit_name = self._ecs2_unit_name(source_unit)
                target_unit_id = self._target_unit_id(source_unit, target_unit_name)
                identity = self._adjacent_property_identity(props, index)
                identity_value = self._resolve_scalar_entry(self._get(identity, 'value')) if identity else None
                explicit_id = self._get(identity_value, 'entryId') if identity else None
                ecs2_property_id = explicit_id or STANDARD_PROPERTY_TO_ECS2.get(property_id) or self._ecs2_id(property_id)
                identity_label = self._get(identity_value, 'label') or self._category_scalar(identity_value)
                if identity_label == explicit_id:
                    identity_label = None
                name = identity_label or STANDARD_PROPERTY_NAMES.get(property_id, self._property_registry_name(property_id))
                metadata_name, defining = self._property_source_metadata(exchange, prop, name)
                name = identity_label or metadata_name or name

                item = {
                    'propertyId': ecs2_property_id,
                    'amount': Decimal(str(self._get(value, 'amount'))),
                    'unitId': target_unit_id,
                    'unitName': self._text(target_unit_name),
                    'name': self._text(name),
                }
                entry_key = ('property', self._get(prop, 'id'))
                uncertainty = self._uncertainty(
                    value,
                    self._entry_metadata_props(entry_key, self._props(prop)),
                    Decimal('1'),
                    'property',
                )
                if uncertainty:
                    item['uncertainty'] = uncertainty
                variable = self._entry_variables.get(entry_key)
                if variable:
                    item['variableName'] = variable
                self._entry_formula(entry_key)
                comments = self._comment_texts(self._get(prop, 'comments'))
                if comments:
                    item['comment'] = self._texts(comments)
                if defining is not None:
                    item['isDefiningValue'] = bool(defining)
                result.append(item)
                if not explicit_id and property_id not in STANDARD_PROPERTY_TO_ECS2:
                    self._master('property', {
                        'id': ecs2_property_id,
                        'name': item['name'],
                        'unitId': target_unit_id,
                        'unitName': item['unitName'],
                    })
        return result

    def _property_source_metadata(self, exchange, prop, name):
        exchange_id = self._get(exchange, 'id')
        candidates = self._property_metadata_cache.get(exchange_id)
        if candidates is None:
            candidates = []
            for metadata in self._props(exchange):
                registry_name = self._property_registry_name(self._get(metadata, 'property'))
                match = re.fullmatch(r'ECS2 property (.+) isDefiningValue', registry_name, re.I)
                if match:
                    candidates.append((match.group(1), self._category_scalar(self._get(metadata, 'value'))))
            self._property_metadata_cache[exchange_id] = candidates

        if not candidates:
            return None, None

        names = [str(name)]
        property_id = self._get(prop, 'property')
        substance = self._get(prop, 'substance')
        if substance:
            names.insert(0, f'{substance} content')
        density_basis = self._get(prop, 'densityBasis')
        if density_basis and density_basis != 'unspecified':
            names.insert(0, f'{density_basis} density')
        if property_id == '5ac08789-7f70-439a-b60e-18c0d612d4ac':
            measure = str(self._get(prop, 'energyMeasure') or '').lower()
            names.extend(['energy content', 'heating value, gross' if 'gross' in measure else 'heating value, net'])
        elif property_id == '1eabd6d5-adc3-4e72-8430-a9679630abf8':
            names.extend(['molar mass', 'molecular weight'])
        elif property_id == '2c74cc35-dcef-4f9e-bda0-89d275526454':
            names.extend(['price', 'market price'])

        normalized = {value.lower().strip() for value in names if value}
        for source_name, value in candidates:
            if source_name.lower().strip() in normalized:
                return source_name, value

        if property_id == '5ac08789-7f70-439a-b60e-18c0d612d4ac':
            for source_name, value in candidates:
                lower = source_name.lower()
                if 'heating value' in lower or 'calorific value' in lower or lower == 'energy content':
                    return source_name, value
        return None, None

    def _adjacent_property_identity(self, props, index):
        if index + 1 >= len(props):
            return None
        candidate = props[index + 1]
        return candidate if self._get(candidate, 'property') == PROPERTY_IDENTITY else None

    def _elementary_master(self, target_flow, unit_id, unit_name, compartment, flow):
        master = {
            'id': target_flow['id'],
            'unitId': unit_id,
            'name': self._text(target_flow['name']),
            'unitName': self._text(unit_name),
            'compartment': compartment,
        }
        cas = self._category_scalar(self._property_value(self._props(flow), CAS_NUMBER))
        formula = self._category_scalar(self._property_value(self._props(flow), CHEMICAL_FORMULA))
        if cas:
            master['casNumber'] = cas
        if formula:
            master['formula'] = formula
        comments = self._comment_texts(self._get(flow, 'comments'))
        if comments:
            master['comment'] = self._texts(comments)
        if self._get(flow, 'sourceIds'):
            self._loss_once(('flow', self._get(flow, 'id'), 'sources'), 'flow_sources_not_representable', f'Flow {self._get(flow, "id")} source references have no EcoSpold2 flow-master source field', len(self._get(flow, 'sourceIds') or []))
        self._master('elementaryExchange', master)

    def _intermediate_master(self, target_flow, unit_id, unit_name, flow, classifications):
        master = {
            'id': target_flow['id'],
            'unitId': unit_id,
            'name': self._text(target_flow['name']),
            'unitName': self._text(unit_name),
        }
        cas = self._category_scalar(self._property_value(self._props(flow), CAS_NUMBER))
        if cas:
            master['casNumber'] = cas
        if classifications:
            master['classification'] = classifications
        properties = self._exchange_properties(flow, {'properties': []}, Decimal('1'))
        if properties:
            master['property'] = properties
        comments = self._comment_texts(self._get(flow, 'comments'))
        if comments:
            master['comment'] = self._texts(comments)
        if self._get(flow, 'sourceIds'):
            self._loss_once(('flow', self._get(flow, 'id'), 'sources'), 'flow_sources_not_representable', f'Flow {self._get(flow, "id")} source references have no EcoSpold2 flow-master source field', len(self._get(flow, 'sourceIds') or []))
        self._master('intermediateExchange', master)

    def _ecs2_unit_name(self, unit, requested=None):
        if unit is None:
            return requested or '1'
        short = self._get(unit, 'shortName') or '1'
        if self._get(unit, 'unitType') == 'currency' and self._get(unit, 'year'):
            return f'{short}{self._get(unit, "year")}'
        return requested or short

    def _target_unit_id(self, unit, target_name):
        mapped = self._mapping_lookup(
            self.unit_mapping,
            self._get(unit, 'id'),
            self._get(unit, 'shortName'),
            self._get(unit, 'semanticName'),
            target_name,
        )
        existing = bool(mapped)
        if isinstance(mapped, str):
            target_id = mapped
        elif isinstance(mapped, dict):
            target_id = self._row(mapped, 'unitId', 'id', 'TargetUnitUUID')
            target_name = self._row(mapped, 'unitName', 'name', 'TargetUnit') or target_name
        elif target_name in (self._get(unit, 'shortName'), self._get(unit, 'semanticName')):
            target_id = self._ecs2_id(self._get(unit, 'id'))
            existing = bool(self._identity_entry(self._current_process, ACTIVITY_NAME_IDENTITY))
        else:
            target_id = self._ecs2_id(uuid_from_string(f'canonical UCUM unit:{target_name}'))
        if not existing:
            self._master('units', {'id': target_id, 'name': self._text(target_name)})
        return target_id

    def _unit_factor(self, source_name, target_name, source_unit=None, target_unit=None):
        if not source_name or not target_name or source_name == target_name:
            return Decimal('1')
        source_ref, source_factor = self._unit_reference(source_name, source_unit)
        target_ref, target_factor = self._unit_reference(target_name, target_unit)
        if source_ref == target_ref:
            return source_factor / target_factor
        ucum = self._ucum()
        if ucum is None:
            return None
        try:
            result = ucum.convertUnitTo(str(source_name), 1, str(target_name))
            status = result.get('status') if hasattr(result, 'get') else getattr(result, 'status', None)
            value = result.get('toVal') if hasattr(result, 'get') else getattr(result, 'toVal', None)
            if status in (None, 'succeeded', 'success') and value is not None:
                return Decimal(str(value))
        except Exception:
            return None
        return None

    def _unit_reference(self, name, unit):
        if unit is None:
            return name, Decimal('1')
        return self._get(unit, 'shortName') or name, Decimal('1')

    def _ucum(self):
        if self._ucum_instance is not None:
            return self._ucum_instance
        if self._ucum_failed:
            return None
        try:
            from javascript import require
            self._ucum_instance = require('@lhncbc/ucum-lhc').UcumLhcUtils.getInstance()
            return self._ucum_instance
        except Exception:
            self._ucum_failed = True
            return None

    def _parameters(self, process):
        parameters = []
        seen = set()
        for parameterization in self._current_parameterizations:
            for model in self._get(parameterization, 'models', []) or []:
                for parameter in self._get(model, 'parameters', []) or []:
                    if self._get(parameter, 'parameterType') == 'choiceInput':
                        self._log(
                            'info',
                            'choice_materialized',
                            f'Choice parameter {self._get(parameter, "id")} was omitted after its solved effects were materialized',
                        )
                        continue
                    if self._get(parameter, 'resultType') != 'quantity':
                        value = self._get(parameter, 'resultingEntry') if self._get(parameter, 'parameterType') == 'formula' else self._get(parameter, 'entry')
                        self._loss_once(
                            ('parameter', self._get(parameter, 'id'), 'category'),
                            'category_parameter_not_representable',
                            f'Canonical category parameter {self._get(parameter, "id")} has no EcoSpold2 numeric parameter representation; solved category is retained only through any fields/formulas it resolved',
                            value=self._category_scalar(value),
                        )
                        continue
                    if self._get(parameter, 'parameterType') == 'inputFromEntry':
                        continue
                    if self._get(parameter, 'id') in self._attached_formula_ids:
                        continue
                    converted = self._parameter(parameter)
                    if converted and converted['parameterId'] not in seen:
                        parameters.append(converted)
                        seen.add(converted['parameterId'])
        return parameters

    def _parameter(self, parameter):
        parameter_type = self._get(parameter, 'parameterType')
        value = self._resolve_scalar_entry(
            self._get(parameter, 'entry') if parameter_type == 'simpleInput' else self._get(parameter, 'resultingEntry')
        )
        if self._get(value, 'quantityType') != 'singleQuantity':
            self._log('warning', 'parameter_not_scalar', f'Parameter {self._get(parameter, "id")} has no solved single quantity')
            return None

        identity = self._identity_entry(parameter, PARAMETER_IDENTITY)
        parameter_id = self._identity_id(self._get(parameter, 'id'), identity, 'parameter') if identity else self._ecs2_id(self._get(parameter, 'id'))
        unit = self.units.get(self._get(value, 'unit'))
        unit_name = self._ecs2_unit_name(unit) if unit else None
        unit_id = self._target_unit_id(unit, unit_name) if unit and unit_name else None
        name = self._get(parameter, 'name') or self._get(parameter, 'variableName') or parameter_id
        result = {
            'parameterId': parameter_id,
            'name': self._text(name),
            'amount': Decimal(str(self._get(value, 'amount'))),
        }
        variable = self._parameter_variable(parameter)
        if variable:
            result['variableName'] = variable
        if unit_id:
            result['unitId'] = unit_id
            result['unitName'] = self._text(unit_name)
        comments = self._comment_texts(self._get(parameter, 'comments'))
        comments.extend(self._flattened_category_comments(parameter))
        if comments:
            result['comment'] = self._texts(comments)
        if self._get(parameter, 'sourceIds'):
            self._loss_once(('parameter', self._get(parameter, 'id'), 'sources'), 'parameter_sources_not_representable', f'Parameter {self._get(parameter, "id")} source references have no EcoSpold2 parameter field', len(self._get(parameter, 'sourceIds') or []))
        uncertainty = self._uncertainty(value, self._props(parameter), Decimal('1'), 'parameter')
        if uncertainty:
            result['uncertainty'] = uncertainty

        if parameter_type == 'formula':
            self.log['summary']['formulaMaterializations'] += 1
            self._log('info', 'formula_materialized', f'Formula parameter {self._get(parameter, "id")} was materialized from its solved result', formula=self._get(parameter, 'formula'))

        if not identity:
            self._master('parameter', {
                'id': parameter_id,
                'name': self._text(name),
                **({'defaultVariableName': variable} if variable else {}),
                **({'unitId': unit_id, 'unitName': self._text(unit_name)} if unit_id else {}),
            })
        return result

    def _build_formula_environment(self, process, parameterizations):
        self._entry_variables = {}
        self._entry_formula_parameters = {}
        self._attached_formula_ids = set()
        self._parameter_index = {}
        self._parameter_variable_overrides = {}
        formula_targets = self._formula_targets(process)

        for parameterization in parameterizations:
            dataset_id = self._get(parameterization, 'id')
            for model in self._get(parameterization, 'models', []) or []:
                for parameter in self._get(model, 'parameters', []) or []:
                    parameter_id = self._get(parameter, 'id')
                    self._parameter_index[(dataset_id, parameter_id)] = parameter
                    self._parameter_index.setdefault(parameter_id, parameter)
                    if self._get(parameter, 'parameterType') == 'formula':
                        targets = formula_targets.get(parameter_id, [])
                        for target in targets:
                            if target in self._entry_formula_parameters:
                                self._log('warning', 'multiple_entry_formulas', f'Multiple formula parameters resolve canonical entry {target[1]}; first formula used')
                            else:
                                self._entry_formula_parameters[target] = parameter
                        if targets:
                            self._attached_formula_ids.add(parameter_id)

        for parameterization in parameterizations:
            for model in self._get(parameterization, 'models', []) or []:
                for parameter in self._get(model, 'parameters', []) or []:
                    if self._get(parameter, 'parameterType') != 'inputFromEntry':
                        continue
                    variable = self._get(parameter, 'variableName')
                    reference = self._get(parameter, 'entryId')
                    if variable and self._get(reference, 'targetType') == 'parameter':
                        referenced = self._referenced_parameter(reference)
                        if referenced is not None:
                            self._parameter_variable_overrides.setdefault(self._get(referenced, 'id'), variable)

        for parameterization in parameterizations:
            for model in self._get(parameterization, 'models', []) or []:
                for parameter in self._get(model, 'parameters', []) or []:
                    parameter_id = self._get(parameter, 'id')
                    parameter_type = self._get(parameter, 'parameterType')
                    targets = formula_targets.get(parameter_id, []) if parameter_type == 'formula' else []
                    variable = self._parameter_variable(parameter)
                    if not variable:
                        continue
                    for target in targets:
                        self._entry_variables[target] = variable
                    if parameter_type == 'inputFromEntry':
                        _, resolved_target = self._resolve_entry_reference(self._get(parameter, 'entryId'))
                        if resolved_target and resolved_target[0] in ('exchange', 'property'):
                            self._entry_variables[resolved_target] = variable

    def _parameter_variable(self, parameter):
        if parameter is None:
            return None
        return self._get(parameter, 'variableName') or self._parameter_variable_overrides.get(self._get(parameter, 'id'))

    def _formula_targets(self, process):
        targets = {}

        def add(parameter_id, target):
            if not parameter_id:
                return
            values = targets.setdefault(parameter_id, [])
            if target not in values:
                values.append(target)

        def resulting_target(value, target):
            quantity_type = self._get(value, 'quantityType')
            category_type = self._get(value, 'categoryType')
            if quantity_type != 'resultingQuantity' and category_type != 'resultingCategory':
                return
            reference = self._get(value, 'parameter')
            add(self._reference_entry(reference, 'parameterId'), target)

        for exchange in self._get(process, 'exchanges', []) or []:
            exchange_id = self._get(exchange, 'id')
            self._formula_target(targets, 'exchange', exchange_id)
            resulting_target(self._get(exchange, 'quantity'), ('exchange', exchange_id))
            for prop in self._props(exchange):
                property_id = self._get(prop, 'id')
                self._formula_target(targets, 'property', property_id)
                resulting_target(self._get(prop, 'value'), ('property', property_id))
            flow = self.flows.get(self._get(exchange, 'flowable'))
            for prop in self._props(flow):
                property_id = self._get(prop, 'id')
                self._formula_target(targets, 'property', property_id)
                resulting_target(self._get(prop, 'value'), ('property', property_id))
        for prop in self._props(process):
            property_id = self._get(prop, 'id')
            self._formula_target(targets, 'property', property_id)
            resulting_target(self._get(prop, 'value'), ('property', property_id))
        return targets

    def _formula_target(self, targets, type_, entry_id):
        if not entry_id:
            return
        parameter_id = uuid_from_string(f'ECS2:parameter:{type_}:{entry_id}:formula')
        values = targets.setdefault(parameter_id, [])
        target = (type_, entry_id)
        if target not in values:
            values.append(target)

    def _entry_formula(self, target, factor=Decimal('1')):
        parameter = self._entry_formula_parameters.get(target)
        if parameter is None:
            return None
        self.log['summary']['formulaMaterializations'] += 1
        self._log(
            'info',
            'formula_materialized',
            f'Formula for canonical {target[0]} {target[1]} was materialized from its solved result',
            formula=self._get(parameter, 'formula'),
        )
        return None

    def _resolve_entry_reference(self, reference, seen=None):
        target_type = self._get(reference, 'targetType')
        file_id = self._get(reference, 'fileId')
        if target_type == 'exchange':
            entry_id = self._reference_entry(reference, 'exchangeId')
            process = self.processes.get(file_id) or self._current_process
            for exchange in self._get(process, 'exchanges', []) or []:
                if self._get(exchange, 'id') == entry_id:
                    return self._resolve_scalar_entry(self._get(exchange, 'quantity'), seen), ('exchange', entry_id)
        if target_type == 'property':
            entry_id = self._reference_entry(reference, 'propertyId')
            datasets = [self.processes.get(file_id), self._current_process]
            for dataset in datasets:
                if dataset is None:
                    continue
                for owner in [dataset] + list(self._get(dataset, 'exchanges', []) or []) + list(self.flows.values()):
                    for prop in self._props(owner):
                        if self._get(prop, 'id') == entry_id:
                            return self._resolve_scalar_entry(self._get(prop, 'value'), seen), ('property', entry_id)
        if target_type == 'parameter':
            parameter = self._referenced_parameter(reference)
            if parameter is None:
                return None, None
            parameter_id = self._get(parameter, 'id')
            parameter_key = (self._get(reference, 'fileId'), parameter_id)
            seen = set() if seen is None else seen
            if parameter_key in seen:
                return None, None
            seen.add(parameter_key)
            parameter_type = self._get(parameter, 'parameterType')
            attached = next(
                (target for target, value in getattr(self, '_entry_formula_parameters', {}).items() if value is parameter),
                None,
            )
            if attached:
                return self._resolve_scalar_entry(self._get(parameter, 'resultingEntry'), seen), attached
            if parameter_type == 'inputFromEntry':
                return self._resolve_entry_reference(self._get(parameter, 'entryId'), seen)
            value = self._get(parameter, 'entry') if parameter_type == 'simpleInput' else self._get(parameter, 'resultingEntry')
            return self._resolve_scalar_entry(value, seen), ('parameter', parameter_id)
        return None, None

    def _referenced_parameter(self, reference):
        parameter_id = self._reference_entry(reference, 'parameterId')
        dataset_id = self._get(reference, 'fileId')
        local = getattr(self, '_parameter_index', {})
        return (
            local.get((dataset_id, parameter_id))
            or local.get(parameter_id)
            or self._global_parameter_index.get((dataset_id, parameter_id))
            or self._global_parameter_index.get(parameter_id)
        )

    def _entry_metadata_props(self, target, *groups):
        result = []
        seen = set()
        for props in groups:
            for prop in props or []:
                key = self._get(prop, 'id') or id(prop)
                if key not in seen:
                    result.append(prop)
                    seen.add(key)
        parameter = getattr(self, '_entry_formula_parameters', {}).get(target)
        for prop in self._props(parameter):
            key = self._get(prop, 'id') or id(prop)
            if key not in seen:
                result.append(prop)
                seen.add(key)
        return result

    def _uncertainty(self, quantity, nearby_props, factor, label):
        quantity = self._resolve_scalar_entry(quantity)
        uncertainty = self._get(quantity, 'uncertainty')
        result = self._distribution(uncertainty, factor, label) or self._preserved_uncertainty(nearby_props, factor)
        undefined = uncertainty is not None and not result
        if undefined:
            result = {'undefined': self._undefined_distribution(quantity, uncertainty, factor)}
        pedigree = self._property(nearby_props, PEDIGREE)
        code = self._category_scalar(self._get(pedigree, 'value')) if pedigree else None
        if code:
            code = str(code)
            if len(code) == 5 and all(character in '12345' for character in code):
                result = result or {}
                result['pedigreeMatrix'] = {
                    'reliability': int(code[0]),
                    'completeness': int(code[1]),
                    'temporalCorrelation': int(code[2]),
                    'geographicalCorrelation': int(code[3]),
                    'furtherTechnologyCorrelation': int(code[4]),
                }

        variance = self._property(nearby_props, VARIANCE_PROPERTY)
        variance_value = self._quantity_amount(self._get(variance, 'value')) if variance else None
        if variance_value is None:
            additional = self._property(nearby_props, ADDITIONAL_VARIANCE)
            additional_value = self._quantity_amount(self._get(additional, 'value')) if additional else None
            distribution = next(iter(result.values()), None) if result and len(result) == 1 else None
            total = self._get(distribution, 'varianceWithPedigreeUncertainty') if distribution else None
            if total is not None and additional_value is not None:
                variance_value = Decimal(str(total)) - Decimal(str(additional_value)) * Decimal(str(factor)) * Decimal(str(factor))
        if variance_value is not None and result:
            for distribution in result.values():
                if isinstance(distribution, dict) and 'varianceWithPedigreeUncertainty' in distribution:
                    distribution['variance'] = Decimal(str(variance_value)) * Decimal(str(factor)) * Decimal(str(factor))
                    break

        comments = []
        if undefined:
            comments.append(f'Canonical {label} uncertainty was materialized as EcoSpold2 undefined uncertainty')
        for prop in nearby_props or []:
            name = self._property_registry_name(self._get(prop, 'property')).lower()
            if name.startswith('ecs2 ') and ' uncertainty comment' in name:
                value = self._category_scalar(self._get(prop, 'value'))
                if value not in (None, ''):
                    comments.append(value)
        if comments:
            result = result or {}
            result['comment'] = self._texts(comments)
        return result

    def _undefined_distribution(self, quantity, uncertainty, factor):
        factor = Decimal(str(factor))
        amount = Decimal(str(self._get(quantity, 'amount', 0))) * factor
        values = [Decimal(str(value)) * factor for value in self._get(uncertainty, 'values', []) or []]
        if not values:
            return {'minValue': amount, 'maxValue': amount, 'standardDeviation95': 0}
        minimum = min(values)
        maximum = max(values)
        return {
            'minValue': minimum,
            'maxValue': maximum,
            'standardDeviation95': max(abs(amount - minimum), abs(maximum - amount)),
        }


    def _preserved_uncertainty(self, props, factor):
        values = {}
        type_ = None
        for prop in props or []:
            name = self._property_registry_name(self._get(prop, 'property'))
            match = re.match(r'ECS2 .+ uncertainty (.+)$', name, re.I)
            if not match:
                continue
            field = match.group(1)
            value = self._category_scalar(self._get(prop, 'value'))
            if field.lower() == 'type':
                type_ = str(value)
            else:
                values[field] = value
        if not type_:
            return None
        factor = Decimal(str(factor))
        scale_fields = {'meanValue', 'minValue', 'mostLikelyValue', 'maxValue', 'scale'}
        square_fields = {'variance', 'varianceWithPedigreeUncertainty'}
        data = {}
        for field, value in values.items():
            decimal = self._decimal(value)
            if decimal is not None and field in scale_fields:
                value = decimal * factor
            elif decimal is not None and field in square_fields:
                value = decimal * factor * factor
            data[field] = value
        return {type_: data}

    def _quantity_amount(self, value):
        value = self._resolve_scalar_entry(value)
        if self._get(value, 'quantityType') != 'singleQuantity':
            return None
        return self._get(value, 'amount')

    def _distribution(self, uncertainty, factor, label):
        if not uncertainty:
            return None
        type_ = self._get(uncertainty, 'uncertaintyType')
        factor = Decimal(str(factor))
        factor2 = factor * factor
        if type_ == 'transformedProbabilityDistribution':
            offset = Decimal(str(self._get(uncertainty, 'offset', 0)))
            if offset:
                self._log('warning', 'uncertainty_not_representable', f'{label} transformed uncertainty with an offset is not representable in EcoSpold2')
                return None
            scale = Decimal(str(self._get(uncertainty, 'scale', 1)))
            return self._distribution(self._get(uncertainty, 'distribution'), factor * scale, label)
        if type_ == 'variance':
            variance = self._get(uncertainty, 'variance')
            if variance is None:
                return None
            return {'normal': {'meanValue': 0, 'varianceWithPedigreeUncertainty': Decimal(str(variance)) * factor2}}
        if type_ == 'truncatedProbabilityDistribution':
            distribution = self._get(uncertainty, 'distribution')
            if self._get(distribution, 'distributionName') != 'Gamma':
                self._log('warning', 'uncertainty_not_representable', f'{label} truncated uncertainty is not representable in EcoSpold2')
                return None
            parameters = self._get(distribution, 'parameters') or {}
            scale = self._get(parameters, 'scale')
            if scale is None and self._get(parameters, 'rate') is not None:
                scale = Decimal('1') / Decimal(str(self._get(parameters, 'rate')))
            return {'gamma': {
                'shape': self._get(parameters, 'shape'),
                'scale': Decimal(str(scale)) * abs(factor),
                'minValue': Decimal(str(self._get(uncertainty, 'minimum', 0))) * factor,
            }}
        if type_ != 'probabilityDistribution':
            self._log('warning', 'uncertainty_not_representable', f'{label} uncertainty type {type_} is not representable in EcoSpold2')
            return None

        name = self._get(uncertainty, 'distributionName')
        parameters = self._get(uncertainty, 'parameters') or {}
        if name == 'Normal':
            return {'normal': {
                'meanValue': Decimal(str(self._get(parameters, 'mean', 0))) * factor,
                'varianceWithPedigreeUncertainty': Decimal(str(self._get(parameters, 'variance', 0))) * factor2,
            }}
        if name == 'Lognormal':
            median = Decimal(str(self._get(parameters, 'median', 0))) * factor
            stdev = Decimal(str(self._get(parameters, 'stdevLog', 0)))
            return {'lognormal': {
                'meanValue': median,
                'mu': median.ln() if median > 0 else 0,
                'varianceWithPedigreeUncertainty': stdev * stdev,
            }}
        if name == 'Triangular':
            return {'triangular': {
                'minValue': Decimal(str(self._get(parameters, 'lowerLimit'))) * factor,
                'mostLikelyValue': Decimal(str(self._get(parameters, 'shape'))) * factor,
                'maxValue': Decimal(str(self._get(parameters, 'upperLimit'))) * factor,
            }}
        if name == 'Uniform':
            return {'uniform': {
                'minValue': Decimal(str(self._get(parameters, 'minimum'))) * factor,
                'maxValue': Decimal(str(self._get(parameters, 'maximum'))) * factor,
            }}
        if name == 'Gamma':
            scale = self._get(parameters, 'scale')
            if scale is None and self._get(parameters, 'rate') is not None:
                scale = Decimal('1') / Decimal(str(self._get(parameters, 'rate')))
            return {'gamma': {'shape': self._get(parameters, 'shape'), 'scale': Decimal(str(scale)) * abs(factor), 'minValue': 0}}
        if name == 'Beta':
            minimum = self._get(parameters, 'minimum')
            maximum = self._get(parameters, 'maximum')
            mode = self._get(parameters, 'mode')
            if None not in (minimum, maximum, mode):
                return {'beta': {
                    'minValue': Decimal(str(minimum)) * factor,
                    'mostLikelyValue': Decimal(str(mode)) * factor,
                    'maxValue': Decimal(str(maximum)) * factor,
                }}
        if name == 'Binomial' and factor == 1:
            return {'binomial': {'n': self._get(parameters, 'n'), 'p': self._get(parameters, 'p')}}
        self._log('warning', 'uncertainty_not_representable', f'{label} distribution {name} is not representable in EcoSpold2')
        return None

    def _system_model(self, system):
        if system is not None:
            prop = self._property(self._props(system), ALLOCATION_MODEL)
            value = self._get(prop, 'value')
            entry_id = self._get(value, 'entryId')
            name = self._category_scalar(value)
            if entry_id:
                model = {'id': entry_id, 'name': name or entry_id, 'shortname': name or 'model'}
                return model
        model = {
            'id': self.config.get('default_system_model_id', DEFAULT_SYSTEM_MODEL_ID),
            'name': self.config.get('default_system_model_name', 'Undefined'),
            'shortname': self.config.get('default_system_model_shortname', 'Undefined'),
        }
        if system is not None:
            self._log('warning', 'system_model_defaulted', f'Production system {self._get(system, "id")} has no EcoSpold2 system-model identity')
        return model

    def _modelling(self, system, system_model):
        project_modelling = self._get(self.project, 'modelling') if self.project is not None else None
        representativeness = {
            'systemModelId': system_model['id'],
            'systemModelName': self._text(system_model['name']),
        }
        sampling = self._get(project_modelling, 'samplingComment')
        extrapolations = self._get(project_modelling, 'extrapolationsComment')
        if sampling:
            representativeness['samplingProcedure'] = self._text(sampling)
        if extrapolations:
            representativeness['extrapolations'] = self._text(extrapolations)
        result = {'representativeness': representativeness}
        reviews = self._reviews()
        if reviews:
            result['review'] = reviews
        return result

    def _reviews(self):
        if self.project is None:
            return []
        result = []
        for activity in self._provenance_activities('dataReview'):
            actors = [self.actors.get(id_) for id_ in self._get(activity, 'actors', []) or []]
            actors = [actor for actor in actors if actor]
            if not actors:
                self._log('warning', 'review_without_actor', f'Review activity {self._get(activity, "id")} has no resolvable actor')
                continue
            reviewer = actors[0]
            version = self._version_parts(self._get(activity, 'reviewedVersion') or self._dataset_version(self._current_process))
            review = {
                'reviewerId': self._actor_id(reviewer),
                'reviewerName': self._actor_name(reviewer),
                'reviewerEmail': self._actor_email(reviewer),
                'reviewDate': self._activity_date(activity),
                'reviewedMajorRelease': version[0],
                'reviewedMinorRelease': version[1],
                'reviewedMajorRevision': version[2],
                'reviewedMinorRevision': version[3],
            }
            comments = self._comment_texts(self._get(activity, 'comments'))
            assessment = self._get(activity, 'assessment', []) or []
            if assessment:
                comments.extend(str(self._category_scalar(item)) for item in assessment)
            if comments:
                review['details'] = {'text': self._texts(comments)}
            result.append(review)
            self._person_master(reviewer)
            if len(actors) > 1:
                self._log('warning', 'review_multiple_actors', f'Review {self._get(activity, "id")} has multiple actors; first actor used as EcoSpold2 reviewer')
        return result

    def _administrative(self, process):
        entry_actor = self._provenance_actor(('dataEntry', 'dataConversion', 'dataModification'))
        generator_actor = self._provenance_actor(('dataGeneration', 'dataCollection', 'dataModelling', 'dataPublication'))
        if entry_actor is None:
            entry_actor = self._fallback_actor('data entry')
        if generator_actor is None:
            generator_actor = entry_actor

        self._person_master(entry_actor)
        self._person_master(generator_actor)
        active = self._dynamic_identity(self._props(entry_actor), 'ecoinvent is active author')
        entry = {
            'personId': self._actor_id(entry_actor),
            'personName': self._actor_name(entry_actor),
            'personEmail': self._actor_email(entry_actor),
            'isActiveAuthor': bool(active) if active is not None else True,
        }
        person_context = self._provenance_metadata('dataEntry', 'data entry.personContextId')
        if person_context:
            entry['personContextId'] = person_context
        publication = {
            'personId': self._actor_id(generator_actor),
            'personName': self._actor_name(generator_actor),
            'personEmail': self._actor_email(generator_actor),
            'isCopyrightProtected': False,
        }
        rights = self._get(self.project, 'rights') if self.project is not None else None
        copyright_ = self._get(rights, 'copyright')
        protected = self._get(copyright_, 'isCopyrightProtected')
        if protected is not None:
            publication['isCopyrightProtected'] = bool(protected)
        access = self._get(rights, 'access')
        access_model = self._get(access, 'accessModel')
        if access_model:
            code = ACCESS_TO_ECS2.get(str(access_model).lower())
            if code is not None:
                publication['accessRestrictedTo'] = code
            else:
                self._log('warning', 'access_model_not_mapped', f'Canonical access model {access_model!r} has no EcoSpold2 access code')
        restricted = self._get(access, 'restrictedTo', []) or []
        for actor_id in restricted:
            actor = self.actors.get(actor_id)
            if actor and self._category_scalar(self._get(actor, 'actorType')) == 'organization':
                company_id = self._actor_id(actor, company=True)
                publication.setdefault('companyId', []).append(company_id)
                publication.setdefault('companyCode', []).append(self._company_code(actor))
                self._company_master(actor)

        for field in (
            'personContextId',
            'publishedSourceIdOverwrittenByChild',
            'publishedSourceContextId',
            'companyIdOverwrittenByChild',
            'companyContextId',
        ):
            value = self._provenance_metadata('dataPublication', f'publication.{field}')
            if value is not None:
                publication[field] = self._metadata_value(value)

        source = self._publication_source()
        if source:
            publication.update(source)
            canonical_source = self.sources.get(source.get('publishedSourceId'))
            if canonical_source is None:
                canonical_source = next(
                    (item for item in self.sources.values() if self._source_id(item) == source.get('publishedSourceId')),
                    None,
                )
            published = self._dynamic_identity(self._props(canonical_source), 'ecoinvent data published in') if canonical_source else None
            if published is not None:
                code = DATA_PUBLISHED_IN_REVERSE.get(str(published).lower())
                if code is not None:
                    publication['dataPublishedIn'] = code
        return {
            'dataEntryBy': entry,
            'dataGeneratorAndPublication': publication,
            'fileAttributes': self._file_attributes(process),
        }

    def _publication_source(self):
        source_ids = self._get(self.project, 'sourceIds', []) or [] if self.project is not None else []
        if not source_ids:
            publication = self._provenance_activities('dataPublication')
            if publication:
                source_ids = [self._reference_source(self._get(publication[0], 'referenceToReport'))]
        source_ids = [id_ for id_ in source_ids if id_]
        if not source_ids:
            return None
        if len(source_ids) > 1:
            self._log('warning', 'multiple_publication_sources', f'{len(source_ids)} canonical publication sources present; first used by EcoSpold2 publication fields')
        source = self.sources.get(source_ids[0])
        if source is None:
            return None
        self._source_master(source)
        return {
            'publishedSourceId': self._source_id(source),
            **({'publishedSourceYear': str(self._get(source, 'year'))} if self._get(source, 'year') else {}),
            'publishedSourceFirstAuthor': self._source_first_author(source),
            **({'pageNumbers': self._get(source, 'pages')} if self._get(source, 'pages') else {}),
        }

    def _file_contexts(self, process):
        values = [(process, 'process')]
        if self._current_system is not None:
            values.append((self._current_system, 'production system'))
        values.extend((item, 'parameterization') for item in self._current_parameterizations)
        if self.project is not None:
            values.append((self.project, 'project'))
        return values

    def _timestamp_value(self, value):
        value = self._serial(value)
        if not value:
            return None
        try:
            parsed = datetime.fromisoformat(value.replace('Z', '+00:00'))
            return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
        except (TypeError, ValueError):
            return None

    def _file_attributes(self, process):
        process_info = self._get(process, 'fileInformation')
        version_info = process_info or self._get(self.project, 'fileInformation')
        version = self._file_version_parts(self._get(self._get(version_info, 'versioning'), 'datasetVersion'))

        creations = []
        edits = []
        for dataset, _ in self._file_contexts(process):
            timestamps = self._get(self._get(dataset, 'fileInformation'), 'timestamp')
            creation = self._timestamp_value(self._get(timestamps, 'creation'))
            last_edit = self._timestamp_value(self._get(timestamps, 'lastEdit'))
            if creation:
                creations.append(creation)
            if last_edit:
                edits.append(last_edit)

        now = datetime.now(timezone.utc)
        creation = min(creations) if creations else now
        last_edit = max(edits) if edits else creation
        if last_edit < creation:
            creation = last_edit

        language = self._project_language() or self.language
        result = {
            'majorRelease': version[0],
            'minorRelease': version[1],
            'majorRevision': version[2],
            'minorRevision': version[3],
            'internalSchemaVersion': self.config.get('internal_schema_version', '2.0.10'),
            'defaultLanguage': language,
            'creationTimestamp': creation.isoformat(),
            'lastEditTimestamp': last_edit.isoformat(),
            'fileGenerator': self.config.get('file_generator', 'Lavoisier canonical to EcoSpold2 conversion'),
            'fileTimestamp': last_edit.isoformat(),
            'contextId': self.config.get('context_id', ECOINVENT_CONTEXT_ID),
        }
        return result

    def _provenance_actor(self, activity_types):
        if self.project is None:
            return None
        for type_ in activity_types:
            for activity in self._provenance_activities(type_):
                for actor_id in self._get(activity, 'actors', []) or []:
                    actor = self.actors.get(actor_id)
                    if actor and self._category_scalar(self._get(actor, 'actorType')) in ('individual', 'software'):
                        return actor
        return None

    def _provenance_activities(self, value):
        provenance = self._get(self.project, 'provenance') if self.project is not None else None
        activities = self._get(provenance, 'activities', []) or []
        return [activity for activity in activities if self._category_scalar(self._get(activity, 'activityType')) == value]

    def _provenance_metadata(self, activity_type, name):
        for activity in self._provenance_activities(activity_type):
            for comment in self._get(activity, 'comments', []) or []:
                if self._get(comment, 'name') == name:
                    return self._get(comment, 'text')
        return None

    def _metadata_value(self, value):
        if isinstance(value, str):
            if value.lower() == 'true':
                return True
            if value.lower() == 'false':
                return False
        return value

    def _fallback_actor(self, role):
        actor_id = uuid_from_string(f'canonical ECS2 fallback actor:{role}')
        actor = {
            'id': actor_id,
            'actorType': 'individual',
            'name': f'Unknown {role} actor',
            'emails': ['unknown@example.invalid'],
        }
        self._log('warning', 'fallback_actor', f'No canonical actor available for EcoSpold2 {role}; explicit placeholder used')
        return actor

    def _identifier_value(self, owner, *systems):
        systems = {system.lower() for system in systems}
        for identifier in self._get(owner, 'identifiers', []) or []:
            if str(self._get(identifier, 'system') or '').lower() in systems:
                return self._get(identifier, 'value')
        return None

    def _actor_id(self, actor, company=False):
        systems = ('ecoinvent company id',) if company else ('ecoinvent person id', 'ecoinvent actor id')
        return self._identifier_value(actor, *systems) or self._ecs2_id(self._get(actor, 'id'))

    def _source_id(self, source):
        return self._identifier_value(source, 'ecoinvent source id') or self._ecs2_id(self._get(source, 'id'))

    def _person_master(self, actor):
        if self._identifier_value(actor, 'ecoinvent person id', 'ecoinvent actor id'):
            return
        phones = self._get(actor, 'phones', []) or []
        self._master('person', {
            'id': self._actor_id(actor),
            'name': self._actor_name(actor),
            'email': self._actor_email(actor),
            **({'telephone': str(phones[0])[:40]} if phones else {}),
        })

    def _company_master(self, actor):
        if self._identifier_value(actor, 'ecoinvent company id', 'ecoinvent actor id'):
            return
        self._master('company', {
            'id': self._actor_id(actor, company=True),
            'code': self._company_code(actor),
            'name': self._text(self._get(actor, 'organizationName') or self._get(actor, 'name')),
            **({'website': self._get(actor, 'websites')[0]} if self._get(actor, 'websites') else {}),
        })

    def _company_code(self, actor):
        value = self._identifier_value(actor, 'ecoinvent company code')
        if value:
            return str(value)[:7]
        return re.sub(r'[^A-Za-z0-9]', '', self._get(actor, 'name') or 'ORG')[:7] or 'ORG'

    def _source_master(self, source):
        if self._identifier_value(source, 'ecoinvent source id'):
            return
        authors = self._get(source, 'authors', []) or []
        first = self._get(authors[0], 'name') if authors else 'Unknown'
        additional = ', '.join(self._get(author, 'name') for author in authors[1:] if self._get(author, 'name')) or None
        self._master('source', {
            'id': self._source_id(source),
            'sourceType': self._source_type(self._get(source, 'sourceType')),
            'firstAuthor': str(first)[:40],
            'title': str(self._get(source, 'title') or 'Untitled canonical source')[:255],
            **({'additionalAuthors': additional[:255]} if additional else {}),
            **({'year': str(self._get(source, 'year'))} if self._get(source, 'year') else {}),
            **({'pageNumbers': self._get(source, 'pages')} if self._get(source, 'pages') else {}),
        })

    def _source_type(self, value):
        text = str(value or '').lower()
        if 'article' in text or 'journal' in text:
            return 1
        if 'book' in text:
            return 2
        if 'report' in text or 'thesis' in text:
            return 3
        if 'web' in text:
            return 0
        return 0

    def _first_source(self, source_ids):
        if not source_ids:
            return None
        if len(source_ids) > 1:
            self._log('warning', 'multiple_sources', f'{len(source_ids)} source references present; first used in EcoSpold2 quantitative reference')
        source = self.sources.get(source_ids[0])
        if source is None:
            return None
        self._source_master(source)
        return {
            'sourceId': self._source_id(source),
            **({'sourceYear': str(self._get(source, 'year'))} if self._get(source, 'year') else {}),
            'sourceFirstAuthor': self._source_first_author(source),
        }

    def _source_first_author(self, source):
        authors = self._get(source, 'authors', []) or []
        return str(self._get(authors[0], 'name') if authors else 'Unknown')[:40]

    def _activity_index(self, activity, geography, time_period, system_model):
        process_identity = self._identity_entry(self._current_process, PROCESS_IDENTITY)
        process_entry_id = self._get(self._get(process_identity, 'value'), 'entryId') if process_identity else None
        direct_ecs2_id = self._ecs2_id(self._get(self._current_process, 'id'))
        existing = bool(self._identity_entry(self._current_process, ACTIVITY_NAME_IDENTITY)) and activity['id'] == direct_ecs2_id
        if activity['id'] != process_entry_id and not existing:
            self._master('activityIndexEntry', {
                'id': activity['id'],
                'activityNameId': activity['activityNameId'],
                'geographyId': geography['geographyId'],
                'startDate': time_period['startDate'],
                'endDate': time_period['endDate'],
                'specialActivityType': activity['specialActivityType'],
                'systemModelId': system_model['id'],
            })
        if not self._identity_entry(self._current_process, ACTIVITY_NAME_IDENTITY):
            self._master('activityName', {'id': activity['activityNameId'], 'name': activity['activityName']})

    def _master(self, name, value, key=None):
        key = key if key is not None else self._get(value, 'id') or self._get(value, 'code') or repr(value)
        seen = self._master_seen.setdefault(name, set())
        if key in seen:
            return
        seen.add(key)
        setattr(self._current_target.userMaster, name, value)

    def _identity_entry(self, owner, property_id):
        return self._property(self._props(owner), property_id)

    def _identity_id(self, canonical_id, prop, label):
        value = self._get(prop, 'value')
        explicit = self._get(value, 'entryId')
        reverse = self._ecs2_id(canonical_id)
        if explicit and explicit != reverse:
            self._log('warning', 'identity_mismatch', f'Canonical {label} reverse UUID {reverse} differs from explicit EcoSpold2 identity {explicit}; explicit identity used', canonicalId=canonical_id)
            return explicit
        return explicit or reverse

    def _ecs2_id(self, canonical_id):
        if not canonical_id:
            return canonical_id
        return uuid_from_uuid(str(canonical_id), CANONICAL_KEY, 'flow_conversion_2')

    def _dynamic_identity(self, props, name):
        for prop in props:
            registry = self.property_registry.get(self._get(prop, 'property'))
            if (self._get(registry, 'name') or '').lower() == name.lower():
                value = self._get(prop, 'value')
                return self._get(value, 'entryId') or self._category_scalar(value)
        return None

    def _property(self, props, property_id):
        return next((prop for prop in props or [] if self._get(prop, 'property') == property_id), None)

    def _properties(self, props, property_id):
        return [prop for prop in props or [] if self._get(prop, 'property') == property_id]

    def _property_value(self, props, property_id):
        prop = self._property(props, property_id)
        return self._resolve_scalar_entry(self._get(prop, 'value')) if prop else None

    def _bool_property(self, owner, property_id):
        return bool(self._category_scalar(self._property_value(self._props(owner), property_id)))

    def _property_registry_name(self, property_id):
        registry = self.property_registry.get(property_id)
        return self._get(registry, 'name') or property_id

    def _mapping_lookup(self, mapping, *keys):
        if not mapping:
            return None
        if isinstance(mapping, dict):
            for key in keys:
                if key is None:
                    continue
                if key in mapping:
                    return mapping[key]
                text = str(key)
                if text in mapping:
                    return mapping[text]
                lower = text.lower()
                for candidate, value in mapping.items():
                    if str(candidate).lower() == lower:
                        return value
            return None
        for row in mapping:
            source_values = [
                self._row(row, 'SourceGeographyCode', 'SourceCode', 'source', 'code', 'SourceFlowContext'),
                self._row(row, 'SourceGeographyUUID', 'SourceUUID', 'id'),
            ]
            if any(str(key).lower() == str(value).lower() for key in keys if key is not None for value in source_values if value is not None):
                return row
        return None

    def _reference_file(self, reference):
        if reference is None:
            return None
        return self._get(reference, 'fileId') or self._get(reference, 'id')

    def _reference_entry(self, reference, key):
        if reference is None:
            return None
        return self._get(reference, key)

    def _reference_source(self, reference):
        if reference is None:
            return None
        return self._get(reference, 'sourceId') or self._get(reference, 'id')

    def _resolve_scalar_entry(self, value, seen=None):
        if value is None or isinstance(value, (str, int, float, bool, Decimal)):
            return value

        quantity_type = self._get(value, 'quantityType')
        category_type = self._get(value, 'categoryType')
        if quantity_type == 'singleQuantity' or category_type == 'singleCategory':
            return value
        if quantity_type != 'resultingQuantity' and category_type != 'resultingCategory':
            return value

        reference = self._get(value, 'parameter')
        parameter_id = self._reference_entry(reference, 'parameterId')
        dataset_id = self._get(reference, 'fileId')
        key = (dataset_id, parameter_id)
        seen = set() if seen is None else seen
        if key in seen:
            return None
        seen.add(key)

        parameter = self._referenced_parameter(reference)
        if parameter is None:
            return None

        parameter_type = self._get(parameter, 'parameterType')
        if parameter_type == 'inputFromEntry':
            resolved, _ = self._resolve_entry_reference(self._get(parameter, 'entryId'), seen)
            return self._resolve_scalar_entry(resolved, seen)
        if parameter_type == 'simpleInput':
            return self._resolve_scalar_entry(self._get(parameter, 'entry'), seen)
        if parameter_type == 'formula':
            return self._resolve_scalar_entry(self._get(parameter, 'resultingEntry'), seen)
        return None

    def _is_scalar_category(self, value):
        return isinstance(value, (str, bool)) or self._get(value, 'categoryType') == 'singleCategory'

    def _category_scalar(self, value):
        value = self._resolve_scalar_entry(value)
        if value is None:
            return None
        if isinstance(value, Enum):
            return self._category_scalar(value.value)
        if isinstance(value, (str, int, float, bool, Decimal)):
            return value
        scalar = self._get(value, 'value')
        if scalar is not None:
            return self._category_scalar(scalar)
        return self._get(value, 'label') or self._get(value, 'code')

    def _reverse_enum(self, value, mapping, default):
        if value is None:
            return default
        if isinstance(value, int):
            return value
        return mapping.get(str(value).lower(), default)

    def _props(self, owner):
        return self._get(owner, 'properties', []) or []

    def _get(self, obj, key, default=None):
        if obj is None:
            return default
        if isinstance(obj, dict):
            if key in obj:
                return obj[key]
            if key == 'from' and 'from_' in obj:
                return obj['from_']
            return default
        if hasattr(obj, key):
            return getattr(obj, key)
        if key == 'from' and hasattr(obj, 'from_'):
            return getattr(obj, 'from_')
        return default

    def _serial(self, value):
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        return value

    def _text(self, value, index=0):
        return {'@index': index, '@lang': self.language, '#text': str(value)}

    def _texts(self, values):
        if not isinstance(values, (list, tuple)):
            values = [values]
        return [self._text(value, index) for index, value in enumerate(values) if value not in (None, '')]

    def _text_value(self, value):
        if isinstance(value, dict):
            return value.get('#text') or value.get('text') or str(value)
        return str(value)

    def _comment_texts(self, comments):
        result = []
        for comment in comments or []:
            text = self._get(comment, 'text')
            if text:
                result.append(text)
        return result

    def _actor_name(self, actor):
        return str(self._get(actor, 'name') or 'Unknown')[:40]

    def _actor_email(self, actor):
        emails = self._get(actor, 'emails', []) or []
        if emails:
            value = str(emails[0])
            return (value[7:] if value.lower().startswith('mailto:') else value)[:80]
        self._log('warning', 'actor_email_missing', f'Actor {self._get(actor, "id")} has no email; placeholder used')
        return 'unknown@example.invalid'

    def _project_language(self):
        language = self._get(self.project, 'language') if self.project is not None else None
        return self._get(language, 'value') if language else None

    def _activity_date(self, activity):
        value = self._get(self._get(activity, 'date'), 'start')
        return str(self._serial(value) or datetime.now(timezone.utc).date().isoformat())[:10]

    def _dataset_version(self, dataset):
        return self._get(self._get(self._get(dataset, 'fileInformation'), 'versioning'), 'datasetVersion') or '1.0.0'

    def _file_version_parts(self, value):
        numbers = [int(number) for number in re.findall(r'\d+', str(value or '1.0.0'))]
        if len(numbers) >= 4:
            return tuple((numbers + [0, 0, 0, 0])[:4])
        major, minor, revision = (numbers + [0, 0, 0])[:3]
        major_revision, minor_revision = divmod(revision, 10)
        return major, minor, major_revision, minor_revision

    def _version_parts(self, value):
        numbers = re.findall(r'\d+', str(value or '1.0.0'))
        parts = [int(number) for number in numbers[:4]]
        return (parts + [0, 0, 0, 0])[:4]

    def _decimal(self, value, default=None):
        if value in (None, '', 'n/a', 'N/A'):
            return default
        try:
            return Decimal(str(value).strip())
        except (InvalidOperation, ValueError):
            return default

    def _looks_uuid(self, value):
        return bool(value and re.fullmatch(r'[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}', str(value)))

    def _log(self, level, code, message, **context):
        entry = {'level': level, 'code': code, 'message': message}
        entry.update({key: value for key, value in context.items() if value is not None})
        if self._current_log_key and self._current_log_key in self.log.get('activities', {}):
            self.log['activities'][self._current_log_key]['messages'].append(entry)
        else:
            self.log.setdefault('global', []).append(entry)
        if level == 'warning':
            self.log['summary']['warnings'] += 1

    def _loss_once(self, key, code, message, count=1, **context):
        if key in self._logged_losses:
            return
        self._logged_losses.add(key)
        self.log['summary']['unconverted'] += count
        self._log('warning', code, message, **context)

    def _log_bundle_losses(self):
        if self.project is not None:
            project_id = self._get(self.project, 'id')
            commissioners = self._get(self.project, 'commissioners', []) or []
            if commissioners:
                self._loss_once(('project', project_id, 'commissioners'), 'commissioners_not_representable', f'{len(commissioners)} project commissioner reference(s) have no direct EcoSpold2 field', len(commissioners))
            hashes = self._get(self.project, 'hashes', []) or []
            if hashes:
                self._loss_once(('project', project_id, 'hashes'), 'hashes_not_representable', f'{len(hashes)} canonical file hash/proof entrie(s) have no EcoSpold2 representation', len(hashes))
            properties = self._get(self.project, 'properties', []) or []
            if properties:
                self._loss_once(('project', project_id, 'properties'), 'project_properties_not_representable', f'{len(properties)} project-level canonical propertie(s) have no direct EcoSpold2 target', len(properties))

            rights = self._get(self.project, 'rights') or {}
            copyright_ = self._get(rights, 'copyright')
            if copyright_:
                if self._get(copyright_, 'referenceToCopyrightOwner'):
                    self._loss_once(('project', project_id, 'copyrightOwner'), 'copyright_owner_not_representable', 'Copyright-owner reference has no EcoSpold2 copyright-owner field')
                if self._get(copyright_, 'referenceToCopyrightStatementURI'):
                    self._loss_once(('project', project_id, 'copyrightStatement'), 'copyright_statement_uri_not_representable', 'Copyright statement URI has no EcoSpold2 target; isCopyrightProtected is preserved')
                if self._get(copyright_, 'validPeriod'):
                    self._loss_once(('project', project_id, 'copyrightPeriod'), 'copyright_period_not_representable', 'Copyright valid period has no EcoSpold2 target')
            license_ = self._get(rights, 'license')
            if license_:
                self._loss_once(('project', project_id, 'license'), 'license_not_representable', 'Canonical license type/URI has no direct EcoSpold2 field')
            access = self._get(rights, 'access')
            if access and self._get(access, 'referenceToPolicyURI'):
                self._loss_once(('project', project_id, 'accessPolicy'), 'access_policy_uri_not_representable', 'Access policy URI has no EcoSpold2 target; access restriction code is preserved when mappable')

            documentation = self._get(self._get(self.project, 'modelling'), 'modellingDocumentation', []) or []
            if documentation:
                self._loss_once(('project', project_id, 'modellingDocumentation'), 'modelling_documentation_not_representable', f'{len(documentation)} modelling-documentation source reference(s) have no dedicated EcoSpold2 field', len(documentation))

        for process in self.processes.values():
            process_id = self._get(process, 'id')
            if self._get(process, 'covarianceMatrix') is not None:
                self._loss_once(('process', process_id, 'covariance'), 'covariance_not_representable', f'Process {process_id} covariance matrix has no EcoSpold2 representation')
            source_ids = self._get(process, 'sourceIds', []) or []
            if source_ids:
                self._loss_once(('process', process_id, 'sources'), 'process_sources_not_representable', f'Process {process_id} has {len(source_ids)} dataset-level source reference(s) without a direct EcoSpold2 process-level source field', len(source_ids))

        for system in self.production_systems.values():
            system_id = self._get(system, 'id')
            transformations = self._get(system, 'appliedTransformationSets', []) or []
            if transformations:
                self._loss_once(('system', system_id, 'transformations'), 'transformations_not_representable', f'Production system {system_id} has {len(transformations)} applied transformation set(s) with no direct EcoSpold2 representation', len(transformations))
            scenarios = self._get(system, 'scenarios', []) or []
            if scenarios:
                self._loss_once(('system', system_id, 'scenarios'), 'scenarios_not_representable', f'Production system {system_id} retains {len(scenarios)} canonical scenario definition(s); solved values are exported but scenario alternatives are not', len(scenarios))
            source_ids = self._get(system, 'sourceIds', []) or []
            if source_ids:
                self._loss_once(('system', system_id, 'sources'), 'production_system_sources_not_representable', f'Production system {system_id} has {len(source_ids)} dataset-level source reference(s) without a direct EcoSpold2 field', len(source_ids))
            for reference in self._get(system, 'functionalReference', []) or []:
                quantity = self._decimal(self._get(reference, 'quantity'))
                target = self._get(reference, 'reference')
                process_id = self._reference_file(self._get(target, 'processId'))
                exchange_id = self._reference_entry(self._get(target, 'outputExchangeId'), 'exchangeId')
                process = self.processes.get(process_id)
                exchange = next(
                    (item for item in self._get(process, 'exchanges', []) or [] if self._get(item, 'id') == exchange_id),
                    None,
                )
                exchange_quantity = self._resolve_scalar_entry(self._get(exchange, 'quantity'))
                exchange_amount = self._decimal(self._get(exchange_quantity, 'amount'))
                if quantity is not None and exchange_amount is not None and quantity != exchange_amount:
                    self._loss_once(
                        ('system', system_id, 'functionalQuantity', exchange_id),
                        'functional_reference_quantity_not_representable',
                        f'Functional-reference normalization quantity {quantity} differs from referenced exchange amount {exchange_amount}; EcoSpold2 can only mark the output as reference product',
                    )

        nonprocess_files = []
        nonprocess_files.extend(self.production_systems.values())
        nonprocess_files.extend(self.parameterizations.values())
        if self.project is not None:
            nonprocess_files.append(self.project)
        if any(self._get(item, 'fileInformation') for item in nonprocess_files):
            self._loss_once(('bundle', 'componentVersions'), 'component_versions_collapsed', 'Production-system, parameterization and project file timestamps/URIs are projected into each activity context, but their independent dataset version numbers cannot all be represented by one EcoSpold2 fileAttributes version')

    def _log_orphan_files(self, materializations):
        used_systems = {self._get(system, 'id') for _, system in materializations if system is not None}
        used_parameterizations = {
            self._reference_file(reference)
            for system in self.production_systems.values()
            for reference in self._get(system, 'parameterizations', []) or []
        }
        for id_ in self.production_systems:
            if id_ not in used_systems:
                self._log('warning', 'production_system_not_materialized', f'Production system {id_} produced no EcoSpold2 activity context')
        for id_ in self.parameterizations:
            if id_ not in used_parameterizations:
                self._log('warning', 'orphan_parameterization', f'Parameterization {id_} is not attached to a production system and was not projected into an EcoSpold2 activity')
